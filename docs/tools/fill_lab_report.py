"""Điền mẫu "Báo cáo kết quả thực hành 04 LAB, Công nghệ phần mềm" của giảng viên bằng nội dung ULS hiện tại.

Giữ nguyên bìa và bảng chữ ký của mẫu, dựng phần thân bằng docx_kit theo quy tắc định dạng của nhóm:
chữ Times New Roman cỡ 13 (đề mục lớn cỡ 14 căn giữa), tiêu đề bảng và hình cỡ 12 in đậm căn giữa nằm dưới,
danh sách dùng dấu -, chữ thể hiện mã dùng Consolas.

Số liệu kiểm thử và coverage lấy từ lần chạy CI thành công gần nhất trên main (run #61); số đo mô hình AI lấy từ
GET /api/recommendations/metrics của hệ thống đang chạy; số API đếm từ /api/docs-json. Cập nhật lại khi đổi.
Usage: python fill_lab_report.py <mẫu.docx> <đầu ra.docx>
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from docx_kit import Template, force_font

ROOT = Path(__file__).resolve().parents[2]
UML = ROOT / "docs" / "uml"
ERD = ROOT / "docs" / "erd"
SHOT = ROOT / "docs" / "report" / "screenshots"
TEXT_W = 16.5

REPO_URL = "https://github.com/khoidev2001/university-library-system"
SRS_URL = f"{REPO_URL}/blob/main/docs/srs/SRS.md"

PROJECT_SUMMARY = (
    "Thư viện trường đại học còn quản lý sách, bạn đọc và mượn trả bằng sổ sách hoặc bảng tính nên tra cứu chậm, "
    "tính phạt dễ sai và không có số liệu thống kê. Đề tài xây dựng Hệ thống Quản lý Thư viện Trường Đại học (ULS) "
    "dạng ứng dụng web: quản lý đầu sách và bản sao có mã vạch; cho mượn tại quầy, trả, gia hạn, báo mất và tự tính "
    "phạt theo chính sách cấu hình được; đặt trước sách, thông báo và email nhắc hạn; nộp phạt trực tuyến qua VNPay; "
    "thống kê, xuất báo cáo, in thẻ và nhãn mã vạch. Hai mô hình AI gợi ý sách cá nhân hoá và dự báo nguy cơ trả trễ. "
    "Công nghệ: Next.js 15 với TailwindCSS, NestJS với Prisma và PostgreSQL, FastAPI với scikit-learn, "
    "Docker Compose, GitHub Actions, triển khai trên Vercel, Render và Neon."
)

TEAM = [
    ["Trần Đình Khôi", "Nhóm trưởng, Scrum Master, Backend và DevOps",
     "SRS và Product Backlog (LAB 1), ERD và đặc tả API (LAB 2), Backend NestJS và Prisma (LAB 3), Docker và CI/CD (LAB 4)", "100%"],
    ["Nguyễn Như Quỳnh", "Frontend và AI/ML",
     "Actor và yêu cầu phi chức năng (LAB 1), UML (LAB 2), giao diện Next.js và dịch vụ AI (LAB 3), kiểm thử (LAB 4)", "100%"],
]

ACTORS = [
    "**Reader (Bạn đọc: sinh viên, giảng viên):** đăng ký, đăng nhập; tìm kiếm và xem chi tiết sách; xem sách đang "
    "mượn, lịch sử, tiền phạt; tự gia hạn; đặt trước sách hết bản; nhận thông báo và email nhắc hạn; nộp phạt trực "
    "tuyến qua VNPay; đánh giá sách; nhận gợi ý AI; đổi mật khẩu, sửa hồ sơ. Chỉ truy cập dữ liệu của chính mình.",
    "**Librarian (Thủ thư):** quyền tra cứu như Reader; CRUD đầu sách, tác giả, thể loại, bản sao; cho mượn tại quầy "
    "bằng mã vạch, nhận trả, báo mất, thu phạt tiền mặt; quản lý hàng đợi đặt trước; tạo bạn đọc tại quầy, khoá và mở "
    "thẻ, cấp lại mật khẩu tạm; xem báo cáo, xuất CSV hoặc Excel, in thẻ và nhãn mã vạch, nhập sách từ tệp; thấy dự "
    "báo nguy cơ trả trễ khi cho mượn.",
    "**Admin (Quản trị):** toàn bộ quyền Librarian; tạo tài khoản Librarian và Admin; cấu hình chính sách mượn (số "
    "cuốn, hạn, gia hạn, mức phạt) theo loại bạn đọc; nhập sách hàng loạt theo ISBN; huấn luyện lại mô hình AI.",
    "**Actor phụ (hệ thống ngoài):** AI Service (gợi ý và dự báo), Google Books API (thông tin sách theo ISBN), "
    "Resend (gửi email), VNPay (cổng thanh toán, môi trường sandbox).",
]

BACKEND_BULLETS = [
    "**Công nghệ:** NodeJS NestJS 10, Prisma 6, PostgreSQL 16; xác thực JWT (Passport), phân quyền RBAC 3 vai trò; "
    "Swagger tự sinh tại `/api/docs`.",
    "**Đã hoàn thiện 78 RESTful API** (đếm trong OpenAPI của hệ thống đang chạy) thuộc 18 nhóm: auth, users, policies, "
    "books, authors, categories, copies, loans, fines, ratings, reports, reservations, notifications, payments, "
    "recommendations, exports, internal, health.",
    "**Kiến trúc** Controller, Service, Repository (Prisma); toàn bộ quy tắc nghiệp vụ (BR-01 đến BR-26) nằm ở tầng "
    "Service. Mọi lỗi trả về một định dạng thống nhất `{ statusCode, error, message, path, timestamp }` kèm mã nghiệp "
    "vụ rõ ràng: `LOAN_LIMIT_REACHED`, `HAS_OVERDUE_LOAN`, `UNPAID_FINE`, `COPY_NOT_AVAILABLE`, `RESERVATION_LIMIT`, "
    "`FINE_PAYMENT_PENDING`.",
    "**Chống đua dữ liệu:** mượn, trả, gia hạn, báo mất, đặt trước, thu phạt và thanh toán dùng khoá hàng bạn đọc "
    "(`FOR UPDATE`), cập nhật có điều kiện và `SKIP LOCKED` khi chọn người đứng đầu hàng đợi; bộ kiểm thử đua bắn yêu "
    "cầu song song trên PostgreSQL thật (ví dụ năm bạn đọc mượn cùng một bản sao thì đúng một phiếu thành công).",
    "**Tích hợp ngoài:** Google Books (nhập sách theo ISBN), Resend (email quên mật khẩu, nhắc hạn), VNPay (thanh toán "
    "phạt, ký và kiểm tra HMAC-SHA512, xử lý IPN và return idempotent, chặn thu trùng).",
    "**Bảo mật:** mật khẩu băm bcrypt, JWT hạn 24 giờ và đọc lại tài khoản ở mỗi yêu cầu (khoá thẻ có hiệu lực ngay), "
    "giới hạn tần suất theo IP, kiểm tra đầu vào, chống formula injection khi xuất CSV.",
]

FRONTEND_BULLETS = [
    "**Công nghệ:** Next.js 15 App Router, TypeScript, TailwindCSS 4, TanStack Query, Recharts (nạp theo yêu cầu), "
    "phông Be Vietnam Pro; 29 trang cho ba nhóm người dùng (bạn đọc, thủ thư, quản trị).",
    "**Thiết kế:** giao diện sáng, bo tròn lớn, hoạt ảnh chỉ dùng `transform` và `opacity` để mượt, tôn trọng thiết "
    "lập giảm chuyển động của hệ điều hành.",
    "**SEO và khả năng tiếp cận:** trang công khai dựng phía máy chủ, tiêu đề và mô tả riêng từng trang, dữ liệu có cấu "
    "trúc schema.org, sitemap và robots tự sinh; mỗi trang một thẻ h1, nhãn cho mọi ô nhập, viền focus rõ ràng.",
    "**Responsive:** kiểm ở bề rộng 390, 768 và 1440 px, không cuộn ngang; bảng dữ liệu chuyển thành thẻ dọc trên "
    "điện thoại; thanh điều hướng thu gọn thành nút menu. Quyền thật vẫn do backend quyết định.",
]

SHOTS = [
    ("01-trang-chu.jpg", "Trang chủ: ô tìm kiếm, từ khoá gợi ý và sách còn trên kệ"),
    ("02-tra-cuu-sach.jpg", "Tra cứu sách với bộ lọc thể loại và tình trạng"),
    ("03-chi-tiet-sach.jpg", "Chi tiết sách kèm tình trạng bản sao và điểm đánh giá"),
    ("04-sach-cua-toi.jpg", "Bạn đọc xem sách đang mượn, hạn trả và lịch sử"),
    ("05-goi-y.jpg", "Trang gợi ý sách cho bạn đọc"),
    ("06-quay-muon-tra.jpg", "Quầy mượn trả của thủ thư (quét mã vạch bản sao và mã bạn đọc)"),
    ("07-thong-ke.jpg", "Thống kê thư viện với biểu đồ lượt mượn và tồn kho"),
    ("14-quan-tri.jpg", "Bảng điều khiển quản trị"),
    ("08-dien-thoai-trang-chu.jpg", "Giao diện trên điện thoại, bề rộng 390 px"),
]

AI_BULLETS = [
    "**Mô hình 1, gợi ý sách (lai):** TF-IDF trên tiêu đề, tác giả, thể loại, mô tả (cosine similarity) trộn 50/50 với "
    "Collaborative Filtering item-based trên ma trận bạn đọc và sách. Đo trên tập giữ lại 20%: Precision@5 = 0,309 và "
    "Recall@5 = 0,110 (1.000 sách, 2.000 bạn đọc, 122.354 tương tác, 300 bạn đọc được đánh giá), vượt ngưỡng nghiệm "
    "thu 0,15.",
    "**Mô hình 2, dự báo nguy cơ trả trễ:** so sánh Logistic Regression với Random Forest trên 10.859 phiếu đã trả "
    "(tỉ lệ trễ 60,1%), tách 20% làm tập kiểm tra. Logistic Regression: Accuracy 0,709, Precision 0,789, Recall 0,703, "
    "F1 0,743, ROC-AUC 0,788. Random Forest: Accuracy 0,719, Precision 0,786, Recall 0,730, F1 0,757, ROC-AUC 0,784. "
    "Hệ thống tự chọn Logistic Regression vì ROC-AUC cao hơn; đạt ngưỡng 0,70.",
    "**Đặc trưng:** số lần mượn trước đó, số lần và tỉ lệ trả trễ trước đó, số sách đang giữ, loại bạn đọc, thể loại "
    "sách, số ngày mượn, tháng mượn và cờ mùa thi; dựng theo trình tự thời gian nên mỗi phiếu chỉ dùng thông tin có "
    "trước ngày mượn.",
    "**Lưu ý trung thực:** lịch sử mượn trả dùng để huấn luyện là dữ liệu mô phỏng sinh từ bộ goodbooks-10k, nên số đo "
    "phản ánh quy luật mô phỏng, chưa phải hành vi bạn đọc thật.",
    "**Chịu lỗi:** khi AI Service không phản hồi, việc cho mượn vẫn thành công với nguy cơ trả trễ rỗng (`null`), trang "
    "gợi ý chuyển sang danh sách sách mượn nhiều, các điểm cuối quản trị AI trả 503 `AI_UNAVAILABLE` thay vì lỗi 500.",
]

AI_SNIPPET = "\n".join([
    "# Snippet 1. ai-service/app/overdue_model.py: chọn mô hình tốt hơn theo ROC-AUC",
    'logistic = Pipeline([("scaler", StandardScaler()),',
    '                     ("clf", LogisticRegression(max_iter=1000, class_weight="balanced"))',
    "                    ]).fit(x_train, y_train)",
    "forest = RandomForestClassifier(n_estimators=200, max_depth=12, min_samples_leaf=5,",
    '                                class_weight="balanced", random_state=42).fit(x_train, y_train)',
    "logistic_scores = score_model(logistic, x_test, y_test)",
    "forest_scores = score_model(forest, x_test, y_test)",
    "chosen_is_forest = forest_scores.roc_auc >= logistic_scores.roc_auc",
    "self.model = forest if chosen_is_forest else logistic",
    "",
    "// Snippet 2. backend/src/modules/ai/ai.client.ts: AI chậm không được chặn quầy mượn (FR-LOAN-02)",
    "const res = await this.fetchImpl(`${this.baseUrl}/predict-overdue`, {",
    "  method: 'POST',",
    "  signal: AbortSignal.timeout(this.predictTimeoutMs),   // 2 giây",
    "});",
    "// lỗi hoặc quá hạn chờ: trả null, phiếu mượn vẫn được tạo bình thường",
])

COMPOSE_INTRO = (
    "Toàn bộ hệ thống khởi chạy bằng một lệnh `docker compose up --build`, gồm 5 container: postgres (cơ sở dữ "
    "liệu), backend (NestJS), ai-service (FastAPI), frontend (Next.js) và gateway (Nginx, cổng 8080). Backend tự áp "
    "migration khi khởi động; frontend, backend và AI Service đều có Dockerfile riêng. Job compose-smoke trong CI "
    "chạy đúng lệnh này trên máy ảo sạch rồi gọi thử ba địa chỉ `/health` qua gateway. Nội dung docker-compose.yml:"
)

TEST_BULLETS = [
    "**Tổng số Test Cases: 576, tất cả đều đạt (100%)** ở lần chạy GitHub Actions run #61 trên main (19/10/2026): "
    "322 unit test backend (Jest, 33 suite), 228 e2e test backend chạy trên PostgreSQL thật (20 suite, gồm các ca đua "
    "song song) và 26 test AI Service (pytest).",
    "**Độ bao phủ mã nguồn (Code Coverage):** backend (unit test, đo trên lớp nghiệp vụ) đạt 91,72% câu lệnh, 76,13% "
    "nhánh, 95,26% hàm, 92,34% dòng; AI Service đạt 93,59%. Ngưỡng được ép trong cấu hình: backend 80% dòng và 75% "
    "nhánh (năm service nặng SQL như thanh toán, xuất dữ liệu, nhập sách, email nhắc hạn, đặt lại mật khẩu được kiểm "
    "bằng e2e riêng nên không tính vào thước đo unit), AI Service 80%; dưới ngưỡng thì pipeline đỏ.",
]

CI_DESCRIPTION = (
    "Workflow `.github/workflows/ci-cd.yml` tự chạy mỗi khi push hoặc mở Pull Request, gồm 4 job song song và 1 job "
    "triển khai. **backend:** lint, build, unit test kèm ngưỡng coverage, e2e test trên PostgreSQL thật (service "
    "container). **ai-service:** pytest kèm ngưỡng coverage. **frontend:** lint và build. **compose-smoke:** dựng cả 5 "
    "container trên máy ảo sạch và gọi thử qua gateway. **deploy:** chỉ chạy khi gộp vào main và mọi job trước đó "
    "xanh, gọi Deploy Hook của Render rồi kiểm tra các địa chỉ trực tuyến."
)

DEPLOY_INTRO = (
    "Kiến trúc triển khai gồm Vercel (frontend), Render (backend và AI Service, Docker) và Neon (PostgreSQL), đều dùng "
    "gói miễn phí. Cấu hình nằm trong repo: `render.yaml`, `frontend/next.config.ts` (chuyển tiếp `/api` và `/ai` sang "
    "Render để trình duyệt chỉ nói chuyện với một địa chỉ) và hướng dẫn từng bước `docs/deploy/DEPLOY.md`. Hệ thống "
    "đã chạy thật trên cloud. Gói miễn phí tự ngủ khi không có truy cập nên lần mở đầu có thể mất khoảng một phút. "
    "Địa chỉ truy cập ứng dụng trực tuyến:"
)

RESULTS = [
    "Hoàn thiện cả 04 bài thực hành LAB theo tiến trình Scrum: Product Backlog 37 User Story (155 điểm), gồm 20 story "
    "MVP trong 4 sprint và 17 story mở rộng ở sprint 5 (đặt trước, thông báo và email nhắc hạn, tài khoản và mật khẩu, "
    "thanh toán VNPay, xuất báo cáo và in thẻ, nhập sách từ tệp); tất cả chạy được trên hệ thống thật.",
    "Hệ thống gồm 78 RESTful API, 29 trang giao diện responsive chuẩn SEO, hai mô hình AI đạt ngưỡng nghiệm thu "
    "(Precision@5 = 0,309 so với ngưỡng 0,15; ROC-AUC = 0,788 so với ngưỡng 0,70), 576 test tự động đạt 100% trên CI.",
    "Áp dụng quy trình DevOps tự động hoá: một lệnh `docker compose up` dựng toàn bộ hệ thống, CI chạy đủ kiểm thử "
    "trên mỗi Pull Request và chỉ triển khai khi xanh, tài liệu thiết kế (ERD, đặc tả API) sinh từ mã nguồn nên luôn "
    "khớp với hệ thống.",
]

LIMITATIONS = (
    "**Hạn chế:** (1) Lịch sử mượn trả dùng để huấn luyện AI là dữ liệu mô phỏng nên độ chính xác phản ánh quy luật "
    "mô phỏng. (2) Sách mẫu lấy từ goodbooks-10k nên tên sách tiếng Anh và chưa có mô tả; nhập sách theo ISBN từ Google "
    "Books đã có mã nguồn và unit test nhưng chưa thử với dữ liệu thật do giới hạn tần suất. (3) Frontend chưa có test "
    "tự động, mới kiểm bằng lint, build và kiểm tay. (4) Thanh toán VNPay mới ở môi trường sandbox; email chỉ gửi tới "
    "được các địa chỉ đã xác minh trên Resend vì chưa có tên miền riêng. (5) JWT hạn một ngày, chưa có refresh token. "
    "(6) Hạ tầng miễn phí không có SLA: Render ngủ sau 15 phút không dùng nên lần gọi đầu mất khoảng một phút."
)

FUTURE = (
    "**Hướng phát triển:** thay dữ liệu mô phỏng bằng dữ liệu mượn trả thật của thư viện rồi huấn luyện lại, so sánh "
    "thêm gradient boosting; xác thực tên miền gửi email và chuyển VNPay sang môi trường thật; bổ sung test giao diện "
    "bằng Playwright và kiểm thử tải để đo chỉ tiêu 2 giây ở p95; thêm nhật ký kiểm toán, cơ chế đồng ý và quyền của "
    "chủ thể dữ liệu cá nhân; lưu mô hình ra bộ nhớ ngoài để AI Service khởi động nhanh sau khi thức dậy."
)


# --------------------------------------------------------------------------- Product Backlog từ SRS
def load_backlog() -> list[list[str]]:
    lines = (ROOT / "docs" / "srs" / "SRS.md").read_text(encoding="utf-8").splitlines()
    start = next(i for i, ln in enumerate(lines) if ln.startswith("| ID | User Story |"))
    rows = []
    for ln in lines[start + 2 :]:
        if not ln.startswith("| US"):
            break
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        ident, story, ac, points = cells[0], cells[1], cells[2], cells[3]
        story = story.replace("**", "")
        items = re.split(r"(?<=[.)])\s+(?=\d+\.\s)", ac)
        ac_short = " ".join(items[:3])
        rows.append([ident, story, ac_short, f"{points} pt"])
    return rows


def set_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_cell(cell, text: str) -> None:
    lines = text.split("\n")
    set_text(cell.paragraphs[0], lines[0])
    for extra in cell.paragraphs[1:]:
        set_text(extra, "")
    for line in lines[1:]:  # dòng thêm nhân bản định dạng của dòng đầu
        clone = copy.deepcopy(cell.paragraphs[0]._p)
        cell._tc.append(clone)
        set_text(Paragraph(clone, cell), line)


def build(template: Path, output: Path) -> None:
    tpl = Template(template)
    tpl.default_half_pt = 26   # mọi chữ cỡ 13
    tpl.force_font_pt = 13     # cả chữ trong bảng
    tpl.center_all_cells = True
    dash_list = tpl.make_dash_list()
    doc = tpl.doc
    P = doc.paragraphs
    style_name = "Table Grid"
    n = {"t": 0, "f": 0}

    # ------------------------------------------------------------------ bìa
    set_text(P[3], "Đề tài : Đề tài 8: Phần mềm Quản lý Thư viện Trường Đại học (ULS)")
    cover = doc.tables[0]
    set_cell(cover.cell(0, 1), "Nhóm 3 - Lớp D23CTC2")
    set_cell(cover.cell(2, 1), "1. Trần Đình Khôi - MSSV: 23Q74802012006 (Nhóm trưởng)\n2. Nguyễn Như Quỳnh - MSSV: 21DQ4802011003")
    set_cell(cover.cell(3, 1), f"GitHub: {REPO_URL}")
    set_cell(cover.cell(4, 1), "Demo Web: https://muce-library.vercel.app")
    set_cell(cover.cell(5, 1), "Ngày 19 tháng 10 năm 2026")

    # ------------------------------------------------------------------ đoạn mẫu rồi bỏ thân mẫu
    tpl.capture("h1", P[5])
    tpl.capture("h2", P[6])          # đã in đậm nghiêng ("1.1.")
    tpl.capture("body", P[7])
    tpl.capture("signdate", next(p for p in P if p.text.startswith("..., ngày")))
    sig_table = copy.deepcopy(doc.tables[5]._tbl)
    cover_table_el = cover._tbl
    node = cover_table_el.getnext()
    while node is not None and node.tag != qn("w:sectPr"):
        nxt = node.getnext()
        tpl.body.remove(node)
        node = nxt

    # ------------------------------------------------------------------ tiện ích dựng thân bài
    def h1(text, new_page=False):
        tpl.add("h1", text, size_half_pt=28, align="center", page_break_before=new_page)

    def h2(text):
        tpl.add("h2", text)

    def body(text, **kw):
        tpl.add("body", text, **kw)

    def bullet(text):
        tpl.add("body", text, list_num=dash_list, list_level=0)

    def last_element():
        return tpl.body.find(qn("w:sectPr")).getprevious()

    def style_caption(el):
        from docx.oxml import OxmlElement
        from docx_kit import _insert_ppr_child, _set_flag, _set_val

        ppr = el.find(qn("w:pPr"))
        for tag in ("w:ind", "w:jc", "w:spacing"):
            for old in ppr.findall(qn(tag)):
                ppr.remove(old)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "60")
        spacing.set(qn("w:after"), "200")
        _insert_ppr_child(ppr, spacing)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        _insert_ppr_child(ppr, jc)
        for r in el.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            _set_flag(rpr, "w:b", True)
            _set_flag(rpr, "w:bCs", True)
            _set_flag(rpr, "w:i", False)
            _set_flag(rpr, "w:iCs", False)
            _set_val(rpr, "w:sz", 24)
            _set_val(rpr, "w:szCs", 24)

    def table(title, header, rows, widths, **kw):
        n["t"] += 1
        tpl.table(header, rows, widths, style=style_name, header_fill="E7E9F0", gap=False, keep_last=True, **kw)
        tpl.add("body", f"Bảng {n['t']}. {title}")
        style_caption(last_element())

    def fig(path, title, w=TEXT_W, h=19.0):
        n["f"] += 1
        tpl.picture(path, w, f"Hình {n['f']}. {title}", caption_kind="body", max_h_cm=h)
        style_caption(last_element())

    # =================================================================== PHẦN I
    h1("PHẦN I: TỔNG QUAN DỰ ÁN VÀ PHÂN CÔNG TƯƠNG TÁC SCRUM")
    h2("1.1. Mô tả tóm tắt dự án")
    body(PROJECT_SUMMARY)
    h2("1.2. Bảng phân công nhiệm vụ và mức độ hoàn thành")
    table("Phân công nhiệm vụ và mức độ hoàn thành", ["Họ và tên SV", "Vai trò trong nhóm", "Nhiệm vụ phụ trách (LAB 1 - 4)", "% Hoàn thành"],
          TEAM, [3.3, 3.5, 7.7, 2.0])

    # =================================================================== PHẦN II
    h1("PHẦN II: KẾT QUẢ THỰC HÀNH TỪNG BÀI LAB", new_page=True)
    h2("2.1. LAB 01: Phân tích yêu cầu & Thiết kế Đặc tả Phần mềm (SRS)")
    body("**a) Xác định Actor:** Danh sách Actor trong hệ thống:")
    for a in ACTORS:
        bullet(a)
    body("**b) Product Backlog & User Stories:** Trích xuất danh sách Product Backlog chính (37 User Story, 155 điểm; "
         "US01 đến US20 là MVP, US21 đến US37 là các đợt mở rộng; Acceptance Criteria rút gọn, đầy đủ trong SRS):")
    table("Product Backlog và User Stories", ["ID", "User Story", "Acceptance Criteria (Điều kiện nghiệm thu)", "Points"],
          load_backlog(), [1.8, 5.5, 7.2, 2.0], center_cols={0, 3})
    body(f"**c) Tài liệu SRS:** Link đường dẫn tới file tài liệu đặc tả SRS đầy đủ chuẩn IEEE 830 (phiên bản 2.0): {SRS_URL}")

    h2("2.2. LAB 02: Thiết kế Kiến trúc Hệ thống & Mô hình hoá UML")
    body("**a) Biểu đồ Use Case Tổng Quan:** ba actor chính (Reader, Librarian, Admin) với quyền lồng nhau Admin gồm "
         "quyền Librarian, Librarian gồm quyền Reader; bốn actor phụ là AI Service, Google Books, Resend và VNPay. "
         "Luồng chính: thủ thư quét mã vạch bản sao và mã bạn đọc, hệ thống kiểm tra chính sách mượn (BR-01 đến BR-05), "
         "gọi AI dự báo trả trễ (include) rồi tạo phiếu; khi nhận trả trễ, hệ thống tự tạo khoản phạt (extend).")
    fig(UML / "use-case-overview.png", "Use Case tổng quan của hệ thống ULS", w=9.5, h=19.5)
    body("**b) Biểu đồ Lớp (Class Diagram):** gói Domain gồm các entity ánh xạ một-một với Prisma schema (User, "
         "LoanPolicy, Book, Author, Category, BookCopy, Loan, Fine, Rating, Reservation, Notification, Payment và "
         "các enum); gói Application gồm các service NestJS và các client ngoài (AiClient, GoogleBooksClient, "
         "MailSender, VNPay). `LoansService` là trung tâm nghiệp vụ, phụ thuộc `PoliciesService`, `FinesService`, "
         "`CopiesService`, `ReservationsService` và `AiClient`.")
    fig(UML / "class-diagram.png", "Biểu đồ lớp của hệ thống", w=TEXT_W, h=19.5)
    body("**c) Biểu đồ Tuần tự (Sequence Diagram):** luồng quan trọng nhất là cho mượn sách tại quầy kèm dự báo AI "
         "(US07, US19): 5 điều kiện chặn BR-01 đến BR-05, gọi AI Service với thời hạn chờ 2 giây (dự phòng risk = null), "
         "một giao dịch tạo phiếu và đổi trạng thái bản sao. Các luồng còn lại: nhận trả, tính phạt và thu tiền; gợi ý "
         "sách; đặt trước sách; nộp phạt trực tuyến qua VNPay.")
    fig(UML / "sequence-borrow.png", "Sequence: cho mượn sách tại quầy kèm dự báo AI (US07, US19)", w=TEXT_W, h=19.5)
    fig(UML / "sequence-return-fine.png", "Sequence: nhận trả sách, tính phạt, thu tiền (US08, US12)", w=TEXT_W, h=19.5)
    fig(UML / "sequence-recommend.png", "Sequence: gợi ý sách cá nhân hoá (US18)", w=TEXT_W, h=19.5)
    fig(UML / "sequence-reservation.png", "Sequence: đặt trước sách và giữ bản sao khi có người trả (US21, US24)", w=TEXT_W, h=19.5)
    fig(UML / "sequence-vnpay-payment.png", "Sequence: nộp phạt trực tuyến qua VNPay (US30)", w=TEXT_W, h=19.5)
    fig(UML / "activity-borrow.png", "Activity Diagram: quy trình cho mượn sách tại quầy", w=12.0, h=19.5)
    body("**d) Thiết kế Cơ sở dữ liệu (Database ERD):** ERD sinh tự động từ Prisma schema, 16 bảng chuẩn 3NF. Bảng "
         "chính: `users` (tài khoản 3 vai trò), `loan_policies` (chính sách mượn theo loại bạn đọc), `books`, `authors`, "
         "`categories`, `book_authors` (quan hệ n-n), `book_copies` (bản sao có mã vạch, trạng thái, gồm HELD khi giữ "
         "cho người đặt trước), `loans` (phiếu mượn, hạn trả, nguy cơ trả trễ), `fines` (phạt OVERDUE và LOST), "
         "`ratings`, `reservations` (đặt trước), `notifications`, `reminder_emails`, `password_reset_tokens`, "
         "`payments` và `payment_fines` (thanh toán VNPay). DDL đầy đủ: `docs/erd/ddl.sql`; sơ đồ kiến trúc và đặc tả "
         "API (Method, URL, Request Body, Response Code, mã lỗi) trong `docs/design/SDD.md`.")
    fig(ERD / "erd.png", "ERD của cơ sở dữ liệu (16 bảng)", w=TEXT_W, h=19.5)
    fig(UML / "architecture.png", "Sơ đồ kiến trúc: gateway, các dịch vụ, cơ sở dữ liệu, tích hợp ngoài", w=TEXT_W, h=19.5)

    h2("2.3. LAB 03: Phát triển Chức năng Cốt lõi & Tích hợp AI/ML")
    body("**a) Kết quả phát triển Backend API:** Mô tả công nghệ sử dụng và các Endpoint chính đã hoàn thiện.")
    for b in BACKEND_BULLETS:
        bullet(b)
    body("**b) Kết quả giao diện Frontend UI/UX:** ảnh chụp từ hệ thống đang chạy tại https://muce-library.vercel.app.")
    for b in FRONTEND_BULLETS:
        bullet(b)
    for filename, caption in SHOTS:
        fig(SHOT / filename, caption, w=TEXT_W if filename != "08-dien-thoai-trang-chu.jpg" else 6.0, h=15.0)
    body("**c) Module AI/ML Tích hợp:** Mô tả thuật toán và kết quả thực nghiệm (số đo lấy từ `GET /api/recommendations/metrics` "
         "của hệ thống đang chạy, huấn luyện ngày 19/10/2026).")
    for b in AI_BULLETS:
        bullet(b)
    tpl.table(["Snippet Tích hợp AI Model (Python/REST API)"], [[AI_SNIPPET]], [TEXT_W], style=style_name,
                           font_pt=10, gap=False, keep_last=True, code=True)
    n["t"] += 1
    tpl.add("body", f"Bảng {n['t']}. Snippet tích hợp mô hình AI")
    style_caption(last_element())

    h2("2.4. LAB 04: Containerization, CI/CD Pipeline & Triển khai Cloud")
    body("**a) Đóng gói Container Docker:** " + COMPOSE_INTRO)
    compose_lines = [ln for ln in (ROOT / "docker-compose.yml").read_text(encoding="utf-8").splitlines() if not ln.lstrip().startswith("#")]
    tpl.table(["Cấu hình docker-compose.yml"], [["\n".join(compose_lines).strip()]], [TEXT_W], style=style_name, font_pt=10,
              gap=False, keep_last=True, code=True)
    n["t"] += 1
    tpl.add("body", f"Bảng {n['t']}. Cấu hình docker-compose.yml")
    style_caption(last_element())
    body("**b) Kết quả Kiểm thử Unit Test:** Kết quả thực thi Unit Test và độ bao phủ mã nguồn (Code Coverage):")
    for b in TEST_BULLETS:
        bullet(b)
    fig(SHOT / "09-coverage-backend.jpg", "Báo cáo coverage của backend trích từ log job backend trên GitHub Actions (run #61)", w=14.0, h=19.0)
    body("**c) Cấu hình CI/CD Pipeline:** " + CI_DESCRIPTION)
    fig(SHOT / "10-ci-pipeline.jpg", "Kết quả pipeline GitHub Actions (run #61 trên main, 4 job đạt và job deploy)", w=TEXT_W, h=19.0)
    body("**d) Triển khai Cloud:** " + DEPLOY_INTRO)
    bullet("URL Web Application: https://muce-library.vercel.app")
    bullet("URL API Swagger Docs: https://muce-library.vercel.app/api/docs (trực tiếp trên Render: https://uls-backend-bo6w.onrender.com/api/docs)")
    fig(SHOT / "11-swagger.jpg", "Swagger UI của hệ thống đang chạy trên cloud", w=TEXT_W, h=15.0)

    # =================================================================== PHẦN III
    h1("PHẦN III: KẾT LUẬN VÀ BÀI HỌC KINH NGHIỆM", new_page=True)
    h2("3.1. Các kết quả chính đã đạt được")
    for r in RESULTS:
        bullet(r)
    h2("3.2. Hạn chế và Hướng phát triển tiếp theo")
    bullet(LIMITATIONS)
    bullet(FUTURE)
    tpl.add("signdate", "Tuy Hòa, ngày 19 tháng 10 năm 2026")
    tpl._append(sig_table)
    tpl.add("body", "")  # đoạn trống cuối

    force_font(doc)
    doc.core_properties.title = "Báo cáo kết quả thực hành 04 LAB Công nghệ phần mềm: ULS"
    doc.core_properties.author = "Trần Đình Khôi, Nguyễn Như Quỳnh"
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: fill_lab_report.py template.docx output.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
    print("wrote", sys.argv[2])
