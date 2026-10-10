"""Hoàn thiện PM_Report.docx: chèn mục lục tự động, số trang; nhờ Word cập nhật trường và xuất PDF.

Usage: python finalize_pm_docx.py <PM_Report.docx> [PM_Report.pdf]
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)


def prepare(path: Path, footer: str = "pm") -> None:
    doc = Document(path)
    for para in doc.paragraphs:
        if para.text.strip() == "[[TOC]]":
            for r in list(para.runs):
                r._r.getparent().remove(r._r)
            add_field(para.add_run(), 'TOC \\o "1-2" \\h \\z \\u')
            break
    for section in doc.sections:
        ft = section.footer
        ft.is_linked_to_previous = False
        p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        if footer == "xy":
            p.add_run("Page ")
            add_field(p.add_run(), "PAGE")
            p.add_run(" / ")
            add_field(p.add_run(), "NUMPAGES")
        else:
            p.add_run("Kế hoạch quản lý dự án ULS   |   Trang ")
            add_field(p.add_run(), "PAGE")
    doc.save(path)


def word_update(docx: Path, pdf: Path) -> int:
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$d = $w.Documents.Open('{docx}')
$d.Repaginate()
foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
$d.Fields.Update() | Out-Null
foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
$d.Repaginate()
$pages = $d.ComputeStatistics(2)
$d.Save()
$d.ExportAsFixedFormat('{pdf}', 17)
$d.Close(0)
$w.Quit()
Write-Output $pages
"""
    out = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=300)
    if out.returncode != 0:
        raise RuntimeError(out.stderr or out.stdout)
    return int(out.stdout.strip().splitlines()[-1])


if __name__ == "__main__":
    docx = Path(sys.argv[1]).resolve()
    pdf = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else docx.with_suffix(".pdf")
    prepare(docx)
    pages = word_update(docx, pdf)
    print(f"{docx.name}: {pages} trang; PDF {pdf}")
