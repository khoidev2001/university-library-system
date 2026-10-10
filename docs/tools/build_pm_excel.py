"""Dựng file Excel kế hoạch quản lý dự án (có công thức sống) từ pm_model.

Usage: python build_pm_excel.py <out.xlsx>
Các sheet dùng tên ASCII để công thức ổn định; tiêu đề tiếng Việt nằm ở ô A1.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

import pm_model as m

HEAD = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="D9E2F3")
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="F2F2F2")
CRIT = PatternFill("solid", fgColor="E06666")
NORM = PatternFill("solid", fgColor="6FA8DC")
GOOD = PatternFill("solid", fgColor="B6D7A8")
WARN = PatternFill("solid", fgColor="FFE599")
BAD = PatternFill("solid", fgColor="EA9999")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
VND = '#,##0'
H1 = '0.0'

wb = Workbook()


def title(ws, text: str, sub: str = "") -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def header(ws, row: int, labels: list[str], start_col: int = 1) -> None:
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.fill = HEAD
        c.font = WHITE_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 32


def widths(ws, ws_widths: list[float]) -> None:
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[col(i)].width = w


def box(ws, rng: str) -> None:
    for row in ws[rng]:
        for c in row:
            c.border = BORDER


# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
title(ws, "Kế hoạch quản lý dự án: Phần mềm Quản lý Thư viện Trường Đại học (Đề tài 8)",
      "Dự án mô phỏng 16 tuần, 4 người. Mọi bảng tính đều dùng công thức; ô vàng là dữ liệu nhập, ô xám là kết quả tính.")
rows = [
    ("GiaDinh", "Tham số dùng chung: đơn giá, năng lực, dự phòng, hệ số 3 điểm, giả định tài chính và EVA."),
    ("WBS", "WBS 3 cấp, ước lượng giờ công 3 điểm (PERT), chi phí kế hoạch (BAC)."),
    ("CPM", "Bảng hoạt động, phụ thuộc FS/SS, ES/EF/LS/LF, độ dự trữ, đường găng."),
    ("PERT", "Ước lượng 3 điểm cho hoạt động và xác suất hoàn thành đúng hạn."),
    ("Gantt", "Biểu đồ Gantt theo tuần; đỏ là hoạt động găng, xanh là hoạt động có độ dự trữ."),
    ("NguonLuc", "Giờ kế hoạch theo tuần và theo vai trò (nền của PV), so với năng lực."),
    ("RuiRo", "Sổ rủi ro: xác suất, tác động, điểm, EMV, ứng phó; ma trận 5x5."),
    ("RACI", "Ma trận trách nhiệm RACI, kiểm tra mỗi gói có đúng một người chịu trách nhiệm cuối (A)."),
    ("TaiChinh", "Chi phí, lợi ích 3 năm, NPV, ROI, thời gian hoàn vốn, độ nhạy."),
    ("EVA", "Phân tích giá trị thu được theo tuần: PV, EV, AC, SV, CV, SPI, CPI, EAC, ETC, VAC, TCPI."),
    ("Timelog", "Giờ công thật ghi nhận của nhóm (từ docs/pm/timelog.csv)."),
    ("ChatLuong", "Mốc chất lượng và tiêu chí chấp nhận, kết quả kiểm thử thật."),
]
header(ws, 4, ["Sheet", "Nội dung"])
for i, (a, b) in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=a).font = BOLD
    ws.cell(row=i, column=2, value=b)
ws["A19"] = "Quy ước màu"
ws["A19"].font = BOLD
for i, (lab, fill) in enumerate([("Ô nhập (đổi được)", INPUT), ("Ô tính (không sửa tay)", CALC), ("Hoạt động găng", CRIT), ("Hoạt động có dự trữ", NORM)], start=20):
    ws.cell(row=i, column=1, value=lab).fill = fill
ws["A25"] = "Lưu ý: EV và AC trong sheet EVA là số mô phỏng có kiểm soát (xem giả định ở sheet GiaDinh), không phải số đo của dự án thật."
ws["A25"].font = Font(italic=True, color="C00000")
widths(ws, [16, 110])

# ---------------------------------------------------------------- GiaDinh
ws = wb.create_sheet("GiaDinh")
title(ws, "Giả định và tham số dùng chung", "Đổi ô vàng thì toàn bộ WBS, PV, EVA và tài chính tự cập nhật.")
header(ws, 4, ["Tham số", "Giá trị", "Ghi chú"])
G: dict[str, str] = {}  # tên -> địa chỉ tuyệt đối


def param(row: int, key: str, label: str, value, note: str = "", fmt: str | None = None) -> None:
    ws.cell(row=row, column=1, value=label)
    c = ws.cell(row=row, column=2, value=value)
    c.fill = INPUT
    if fmt:
        c.number_format = fmt
    ws.cell(row=row, column=3, value=note)
    G[key] = f"GiaDinh!$B${row}"


param(5, "rate", "Đơn giá giờ công (đ/giờ)", m.RATE, "Giả định sinh viên thực tập", VND)
param(6, "weeks", "Số tuần dự án", m.WEEKS, "1 tuần thật = 2 tuần mô phỏng")
param(7, "cap", "Năng lực mỗi người (giờ/tuần)", m.CAPACITY_PER_PERSON, "Bán thời gian, khoảng 2 ngày/tuần")
param(8, "people", "Số người", 4, "PM, Frontend và AI, Backend và DevOps, BA và QA")
param(9, "reserve", "Dự phòng quản lý (% BAC)", m.RESERVE_RATE, "Dùng cho thay đổi được duyệt và rủi ro", "0%")
param(10, "disc", "Tỉ lệ chiết khấu", m.FIN["discount"], "Chi phí vốn giả định", "0%")
ws["A12"] = "Hệ số ước lượng 3 điểm theo mức bất định"
ws["A12"].font = BOLD
header(ws, 13, ["Mức", "Lạc quan (O = M x)", "Bi quan (P = M x)"])
for i, (k, (lo, hi)) in enumerate(m.UNCERTAINTY.items(), start=14):
    ws.cell(row=i, column=1, value=k)
    ws.cell(row=i, column=2, value=lo).fill = INPUT
    ws.cell(row=i, column=3, value=hi).fill = INPUT
UNC_RANGE = "GiaDinh!$A$14:$C$16"
ws["E12"] = "Hệ số 3 điểm cho hoạt động (sheet PERT)"
ws["E12"].font = BOLD
header(ws, 13, ["Mức", "Lạc quan (a = m x)", "Bi quan (b = m x)"], start_col=5)
for i, (k, (lo, hi)) in enumerate(m.ACT_UNCERTAINTY.items(), start=14):
    ws.cell(row=i, column=5, value=k)
    ws.cell(row=i, column=6, value=lo).fill = INPUT
    ws.cell(row=i, column=7, value=hi).fill = INPUT
ACT_UNC_RANGE = "GiaDinh!$E$14:$G$16"

ws["A18"] = "Giả định tài chính"
ws["A18"].font = BOLD
header(ws, 19, ["Tham số", "Giá trị", "Ghi chú"])
fin_rows = [
    ("readers", "Số bạn đọc (sinh viên, giảng viên)", m.FIN["reader_count"], "Giả định quy mô trường", VND),
    ("lpr", "Lượt mượn mỗi bạn đọc mỗi năm", m.FIN["loans_per_reader_year"], "", None),
    ("min_loan", "Phút thủ thư tiết kiệm mỗi lượt mượn trả", m.FIN["minutes_saved_per_loan"], "Ghi sổ tay 3 phút, quét mã 1 phút", None),
    ("lookups", "Lượt tra cứu hộ mỗi năm", m.FIN["lookups_per_year"], "", VND),
    ("min_look", "Phút tiết kiệm mỗi lượt tra cứu", m.FIN["minutes_saved_per_lookup"], "", None),
    ("lib_rate", "Đơn giá giờ thủ thư (đ/giờ)", m.FIN["librarian_rate"], "", VND),
    ("rep_h", "Giờ làm báo cáo thống kê tiết kiệm mỗi năm", m.FIN["report_hours_saved"], "12 báo cáo x 8 giờ", None),
    ("lost_v", "Giá trị sách mất mỗi năm hiện nay (đ)", m.FIN["lost_book_value_year"], "", VND),
    ("lost_r", "Tỉ lệ giảm thất thoát nhờ theo dõi quá hạn", m.FIN["lost_reduction"], "", "0%"),
    ("ad1", "Mức áp dụng năm 1", m.FIN["adoption"][0], "70% do làm quen", "0%"),
    ("ad2", "Mức áp dụng năm 2", m.FIN["adoption"][1], "", "0%"),
    ("ad3", "Mức áp dụng năm 3", m.FIN["adoption"][2], "", "0%"),
    ("host", "Hạ tầng trả phí tối thiểu (đ/năm)", m.FIN["ops_hosting"], "Phương án thay free-tier khi vận hành thật", VND),
    ("mh", "Giờ bảo trì mỗi năm", m.FIN["ops_maintenance_hours"], "", None),
    ("train", "Đào tạo, hỗ trợ mỗi năm (đ)", m.FIN["ops_training"], "", VND),
]
for i, (k, lab, v, note, fmt) in enumerate(fin_rows, start=20):
    param(i, k, lab, v, note, fmt)

ws["A37"] = "Giả định EVA (số mô phỏng có kiểm soát)"
ws["A37"].font = BOLD
header(ws, 38, ["Giai đoạn", "Hiệu suất tiến độ (EV/PV trong tuần)", "Hệ số chi phí (giờ thực/giờ thu được)", "Lý do"])
eva_phase = [("Tuần 1-4", 0.96, 1.04, m.EVA_NOTES["1-4"]), ("Tuần 5-8", 0.80, 1.12, m.EVA_NOTES["5-8"]),
             ("Tuần 9-12", 1.10, 1.10, m.EVA_NOTES["9-12"]), ("Tuần 13-16", 1.12, 1.05, m.EVA_NOTES["13-16"])]
for i, (lab, p, c, why) in enumerate(eva_phase, start=39):
    ws.cell(row=i, column=1, value=lab)
    ws.cell(row=i, column=2, value=p).fill = INPUT
    ws.cell(row=i, column=3, value=c).fill = INPUT
    ws.cell(row=i, column=4, value=why)
EVA_PHASE_ROW = {1: 39, 2: 40, 3: 41, 4: 42}
widths(ws, [46, 22, 26, 90, 8, 18, 18])

# ---------------------------------------------------------------- WBS
ws = wb.create_sheet("WBS")
title(ws, "WBS và ước lượng giờ công (PERT 3 điểm)", "E = (O + 4M + P) / 6; sigma = (P - O) / 6. O và P tính từ M theo mức bất định ở sheet GiaDinh.")
header(ws, 4, ["Mã WBS", "Gói công việc (Work Package)", "Pha", "Vai trò", "Hoạt động", "Mức", "M (giờ)", "O (giờ)", "P (giờ)",
               "E (giờ)", "Sigma", "Chi phí (đ)", "Sản phẩm bàn giao"])
wps = sorted(m.WORK_PACKAGES, key=lambda w: m.wbs_sort_key(w.wbs))
WBS_FIRST = 5
for i, wp in enumerate(wps, start=WBS_FIRST):
    phase = m.WBS_L1[wp.wbs.split(".")[0]]
    ws.cell(row=i, column=1, value=wp.wbs)
    ws.cell(row=i, column=2, value=wp.name)
    ws.cell(row=i, column=3, value=phase)
    ws.cell(row=i, column=4, value=wp.owner)
    ws.cell(row=i, column=5, value=wp.activity)
    ws.cell(row=i, column=6, value=wp.cls).fill = INPUT
    ws.cell(row=i, column=7, value=wp.m).fill = INPUT
    ws.cell(row=i, column=8, value=f"=ROUND(G{i}*VLOOKUP(F{i},{UNC_RANGE},2,FALSE),1)")
    ws.cell(row=i, column=9, value=f"=ROUND(G{i}*VLOOKUP(F{i},{UNC_RANGE},3,FALSE),1)")
    ws.cell(row=i, column=10, value=f"=(H{i}+4*G{i}+I{i})/6")
    ws.cell(row=i, column=11, value=f"=(I{i}-H{i})/6")
    ws.cell(row=i, column=12, value=f"=J{i}*{G['rate']}")
    ws.cell(row=i, column=13, value=wp.deliverable)
    for c in range(8, 13):
        ws.cell(row=i, column=c).fill = CALC
    ws.cell(row=i, column=10).number_format = H1
    ws.cell(row=i, column=11).number_format = '0.00'
    ws.cell(row=i, column=12).number_format = VND
WBS_LAST = WBS_FIRST + len(wps) - 1
tot = WBS_LAST + 1
ws.cell(row=tot, column=2, value="TỔNG (BAC)").font = BOLD
for c, f in ((7, f"=SUM(G{WBS_FIRST}:G{WBS_LAST})"), (10, f"=SUM(J{WBS_FIRST}:J{WBS_LAST})"), (12, f"=SUM(L{WBS_FIRST}:L{WBS_LAST})")):
    cell = ws.cell(row=tot, column=c, value=f)
    cell.font = BOLD
    cell.fill = SUB
cell = ws.cell(row=tot + 1, column=2, value="Dự phòng quản lý")
ws.cell(row=tot + 1, column=12, value=f"=L{tot}*{G['reserve']}").number_format = VND
ws.cell(row=tot + 2, column=2, value="Ngân sách đề xuất (BAC + dự phòng)").font = BOLD
ws.cell(row=tot + 2, column=12, value=f"=L{tot}+L{tot+1}").number_format = VND
ws.cell(row=tot, column=12).number_format = VND
ws.cell(row=tot, column=10).number_format = H1
BAC_CELL = f"WBS!$L${tot}"
BAC_H_CELL = f"WBS!$J${tot}"
RESERVE_TOTAL_CELL = f"WBS!$L${tot+2}"
# tổng theo pha
ws.cell(row=tot + 4, column=2, value="Tổng theo pha WBS cấp 1").font = BOLD
header(ws, tot + 5, ["Mã", "Pha", "", "", "", "", "", "", "", "E (giờ)", "% giờ", "Chi phí (đ)"])
for k, (code, name) in enumerate(m.WBS_L1.items()):
    r = tot + 6 + k
    ws.cell(row=r, column=1, value=code)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=10, value=f'=SUMIF($C${WBS_FIRST}:$C${WBS_LAST},B{r},$J${WBS_FIRST}:$J${WBS_LAST})').number_format = H1
    ws.cell(row=r, column=11, value=f"=J{r}/$J${tot}").number_format = "0.0%"
    ws.cell(row=r, column=12, value=f"=J{r}*{G['rate']}").number_format = VND
widths(ws, [9, 58, 14, 8, 10, 6, 9, 9, 9, 9, 8, 14, 28])
ws.freeze_panes = "C5"

# ---------------------------------------------------------------- CPM
ws = wb.create_sheet("CPM")
title(ws, "Lập lịch và đường găng (CPM)",
      "ES, EF tính xuôi; LS, LF tính ngược; Slack = LS - ES. Quan hệ FS và SS có độ trễ (lag). Đơn vị: tuần.")
header(ws, 4, ["Mã", "Hoạt động", "Thời lượng (tuần)", "Phụ thuộc (tiền nhiệm)", "ES", "EF", "LS", "LF", "Slack", "Găng?"])
ACT_ROW = {a.code: 5 + i for i, a in enumerate(m.ACTIVITIES)}
succ: dict[str, list[tuple[str, str, int]]] = {a.code: [] for a in m.ACTIVITIES}
for a in m.ACTIVITIES:
    for pc, kind, lag in a.preds:
        succ[pc].append((a.code, kind, lag))


def pred_text(a: m.Activity) -> str:
    if not a.preds:
        return "-"
    return ", ".join(f"{pc}({kind}{'+'+str(lag) if lag else ''})" for pc, kind, lag in a.preds)


end_cell = f"CPM!$F${5 + len(m.ACTIVITIES)}"
for a in m.ACTIVITIES:
    r = ACT_ROW[a.code]
    ws.cell(row=r, column=1, value=a.code)
    ws.cell(row=r, column=2, value=a.name)
    ws.cell(row=r, column=3, value=a.dur).fill = INPUT
    ws.cell(row=r, column=4, value=pred_text(a))
    # ES
    terms = ["0"]
    for pc, kind, lag in a.preds:
        pr = ACT_ROW[pc]
        terms.append(f"{'F' if kind == 'FS' else 'E'}{pr}+{lag}")
    ws.cell(row=r, column=5, value="=MAX(" + ",".join(terms) + ")")
    ws.cell(row=r, column=6, value=f"=E{r}+C{r}")
    # LF
    bounds = [f"$F${5 + len(m.ACTIVITIES)}"]
    for sc, kind, lag in succ[a.code]:
        sr = ACT_ROW[sc]
        if kind == "FS":
            bounds.append(f"G{sr}-{lag}")
        else:
            bounds.append(f"G{sr}-{lag}+C{r}")
    ws.cell(row=r, column=8, value="=MIN(" + ",".join(bounds) + ")")
    ws.cell(row=r, column=7, value=f"=H{r}-C{r}")
    ws.cell(row=r, column=9, value=f"=G{r}-E{r}")
    ws.cell(row=r, column=10, value=f'=IF(I{r}=0,"GĂNG","")')
    for c in range(5, 11):
        ws.cell(row=r, column=c).fill = CALC
        ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
last_act = 5 + len(m.ACTIVITIES) - 1
ws.cell(row=last_act + 1, column=2, value="Thời gian dự án (tuần)").font = BOLD
ws.cell(row=last_act + 1, column=6, value=f"=MAX(F5:F{last_act})").font = BOLD
ws.conditional_formatting.add(f"A5:J{last_act}", FormulaRule(formula=[f"$I5=0"], fill=PatternFill("solid", bgColor="F4CCCC")))
box(ws, f"A4:J{last_act}")
widths(ws, [6, 38, 12, 30, 7, 7, 7, 7, 8, 9])
PROJECT_END = f"CPM!$F${last_act + 1}"

ws.cell(row=last_act + 4, column=1, value="Các đường găng (chuỗi hoạt động có Slack 0 nối nhau bằng ràng buộc chặt)").font = BOLD
for k, path in enumerate(m.critical_paths(), start=1):
    ws.cell(row=last_act + 4 + k, column=1, value=f"Đường {k}:")
    ws.cell(row=last_act + 4 + k, column=2, value=" → ".join(path))

# ---------------------------------------------------------------- PERT
ws = wb.create_sheet("PERT")
title(ws, "Ước lượng 3 điểm cho hoạt động và xác suất đúng hạn",
      "E = (a + 4m + b) / 6; Var = ((b - a) / 6)^2. Thời gian lịch kỳ vọng = thời gian CPM + tổng (E - m) trên đường găng chính.")
header(ws, 4, ["Mã", "Hoạt động", "Mức", "a (lạc quan)", "m (khả dĩ)", "b (bi quan)", "E", "Phương sai", "Trên đường găng chính?"])
main = m.critical_chain()
for a in m.ACTIVITIES:
    r = ACT_ROW[a.code]
    ws.cell(row=r, column=1, value=a.code)
    ws.cell(row=r, column=2, value=a.name)
    ws.cell(row=r, column=3, value=a.cls).fill = INPUT
    ws.cell(row=r, column=4, value=f"=ROUND(E{r}*VLOOKUP(C{r},{ACT_UNC_RANGE},2,FALSE),2)")
    ws.cell(row=r, column=5, value=f"=CPM!C{r}")
    ws.cell(row=r, column=6, value=f"=ROUND(E{r}*VLOOKUP(C{r},{ACT_UNC_RANGE},3,FALSE),2)")
    ws.cell(row=r, column=7, value=f"=(D{r}+4*E{r}+F{r})/6")
    ws.cell(row=r, column=8, value=f"=((F{r}-D{r})/6)^2")
    ws.cell(row=r, column=9, value="x" if a.code in main else "")
    for c in (4, 6, 7, 8):
        ws.cell(row=r, column=c).fill = CALC
        ws.cell(row=r, column=c).number_format = '0.00'
sr = last_act + 2
ws.cell(row=sr, column=2, value="Đường găng chính (độ bất định lớn nhất): " + " → ".join(main)).font = BOLD
ws.cell(row=sr + 1, column=2, value="Thời gian CPM (tuần)")
ws.cell(row=sr + 1, column=7, value=f"={PROJECT_END}")
ws.cell(row=sr + 2, column=2, value="Thời gian lịch kỳ vọng (tuần)")
ws.cell(row=sr + 2, column=7, value=f'=G{sr+1}+SUMPRODUCT((I5:I{last_act}="x")*(G5:G{last_act}-E5:E{last_act}))')
ws.cell(row=sr + 3, column=2, value="Độ lệch chuẩn (tuần)")
ws.cell(row=sr + 3, column=7, value=f'=SQRT(SUMPRODUCT((I5:I{last_act}="x")*H5:H{last_act}))')
ws.cell(row=sr + 4, column=2, value="Xác suất hoàn thành trong 16 tuần")
ws.cell(row=sr + 4, column=7, value=f"=NORMSDIST((16-G{sr+2})/G{sr+3})")
ws.cell(row=sr + 5, column=2, value="Xác suất hoàn thành trong 17 tuần")
ws.cell(row=sr + 5, column=7, value=f"=NORMSDIST((17-G{sr+2})/G{sr+3})")
ws.cell(row=sr + 6, column=2, value="Xác suất hoàn thành trong 18 tuần")
ws.cell(row=sr + 6, column=7, value=f"=NORMSDIST((18-G{sr+2})/G{sr+3})")
for k in range(1, 7):
    ws.cell(row=sr + k, column=7).fill = CALC
    ws.cell(row=sr + k, column=7).number_format = '0.00' if k < 4 else '0%'
PERT_P16 = f"PERT!$G${sr+4}"
widths(ws, [6, 38, 6, 12, 12, 12, 9, 12, 22])

# ---------------------------------------------------------------- Gantt
ws = wb.create_sheet("Gantt")
title(ws, "Biểu đồ Gantt theo tuần", "Ô có số 1 nghĩa là hoạt động chạy trong tuần đó. Đỏ: hoạt động găng; xanh: có độ dự trữ.")
header(ws, 4, ["Mã", "Hoạt động", "Bắt đầu (ES)", "Kết thúc (EF)", "Găng?"] + [f"T{w}" for w in range(1, m.WEEKS + 1)])
for a in m.ACTIVITIES:
    r = ACT_ROW[a.code]
    ws.cell(row=r, column=1, value=a.code)
    ws.cell(row=r, column=2, value=a.name)
    ws.cell(row=r, column=3, value=f"=CPM!E{r}")
    ws.cell(row=r, column=4, value=f"=CPM!F{r}")
    ws.cell(row=r, column=5, value=f"=CPM!J{r}")
    for w in range(1, m.WEEKS + 1):
        cl = col(5 + w)
        ws.cell(row=r, column=5 + w, value=f"=IF(AND({w}>$C{r},{w}<=$D{r}),1,0)")
        ws.cell(row=r, column=5 + w).number_format = ';;;'
        ws.column_dimensions[cl].width = 3.6
rng = f"F5:{col(5 + m.WEEKS)}{last_act}"
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(F5=1,$E5="GĂNG")'], fill=PatternFill("solid", bgColor="E06666")))
ws.conditional_formatting.add(rng, FormulaRule(formula=["F5=1"], fill=PatternFill("solid", bgColor="6FA8DC")))
box(ws, f"A4:{col(5 + m.WEEKS)}{last_act}")
widths(ws, [6, 38, 11, 11, 8])
ws.cell(row=last_act + 2, column=2, value="Mốc (milestone)").font = BOLD
milestones = [(2, "M1: Hiến chương, SOW được duyệt"), (3, "M2: SRS và 20 User Story được duyệt"), (5, "M3: Thiết kế (kiến trúc, CSDL, API, giao diện) được duyệt"),
              (9, "M4: Backend và Frontend lõi chạy được (mượn trả đầu-cuối)"), (12, "M5: AI hoàn chỉnh, đạt ngưỡng chỉ số"),
              (14, "M6: Tích hợp xong, bắt đầu kiểm thử hệ thống"), (15, "M7: Qua UAT, triển khai cloud"), (16, "M8: Nghiệm thu và bàn giao")]
for k, (wk, text) in enumerate(milestones, start=1):
    ws.cell(row=last_act + 2 + k, column=1, value=f"T{wk}")
    ws.cell(row=last_act + 2 + k, column=2, value=text)

# ---------------------------------------------------------------- NguonLuc
ws = wb.create_sheet("NguonLuc")
title(ws, "Giờ kế hoạch theo tuần (nền của PV) và tải theo vai trò",
      "Gói công việc trải đều trong cửa sổ [ES, EF] của hoạt động (lịch sớm). Công việc duy trì (LOE) trải đều 16 tuần, riêng 7.3 dồn vào tuần 4, 8, 12, 16.")
header(ws, 4, ["Mã WBS", "Gói công việc", "Vai trò", "Hoạt động", "E (giờ)", "ES", "EF"] + [f"T{w}" for w in range(1, m.WEEKS + 1)])
NL_FIRST = 5
for i, wp in enumerate(wps, start=NL_FIRST):
    wr = WBS_FIRST + (i - NL_FIRST)
    ws.cell(row=i, column=1, value=f"=WBS!A{wr}")
    ws.cell(row=i, column=2, value=f"=WBS!B{wr}")
    ws.cell(row=i, column=3, value=f"=WBS!D{wr}")
    ws.cell(row=i, column=4, value=f"=WBS!E{wr}")
    ws.cell(row=i, column=5, value=f"=WBS!J{wr}").number_format = H1
    if wp.activity == "LOE":
        ws.cell(row=i, column=6, value=0)
        ws.cell(row=i, column=7, value=m.WEEKS)
    else:
        ar = ACT_ROW[wp.activity]
        ws.cell(row=i, column=6, value=f"=CPM!E{ar}")
        ws.cell(row=i, column=7, value=f"=CPM!F{ar}")
    for w in range(1, m.WEEKS + 1):
        c = col(7 + w)
        if wp.wbs == "7.3":
            ws.cell(row=i, column=7 + w, value=f"=IF(OR({w}=4,{w}=8,{w}=12,{w}=16),$E{i}/4,0)")
        else:
            ws.cell(row=i, column=7 + w, value=f"=IF(AND({w}>$F{i},{w}<=$G{i}),$E{i}/($G{i}-$F{i}),0)")
        ws.cell(row=i, column=7 + w).number_format = '0.0'
        ws.column_dimensions[c].width = 6
NL_LAST = NL_FIRST + len(wps) - 1
r0 = NL_LAST + 2
ws.cell(row=r0, column=2, value="Tổng giờ kế hoạch theo tuần").font = BOLD
for w in range(1, m.WEEKS + 1):
    c = col(7 + w)
    ws.cell(row=r0, column=7 + w, value=f"=SUM({c}{NL_FIRST}:{c}{NL_LAST})").number_format = '0.0'
role_row = {}
for k, role in enumerate(m.ROLES, start=1):
    r = r0 + k
    role_row[role] = r
    ws.cell(row=r, column=2, value=f"Vai trò {role}: {m.ROLES[role][0]}")
    ws.cell(row=r, column=3, value=role)
    ws.cell(row=r, column=5, value=f'=SUMIF($C${NL_FIRST}:$C${NL_LAST},C{r},$E${NL_FIRST}:$E${NL_LAST})').number_format = H1
    for w in range(1, m.WEEKS + 1):
        c = col(7 + w)
        ws.cell(row=r, column=7 + w, value=f'=SUMIF($C${NL_FIRST}:$C${NL_LAST},$C{r},{c}${NL_FIRST}:{c}${NL_LAST})').number_format = '0.0'
cap_r = r0 + 5
ws.cell(row=cap_r, column=2, value="Năng lực nhóm (giờ/tuần)")
for w in range(1, m.WEEKS + 1):
    ws.cell(row=cap_r, column=7 + w, value=f"={G['cap']}*{G['people']}")
util_r = cap_r + 1
ws.cell(row=util_r, column=2, value="Mức sử dụng nhóm").font = BOLD
for w in range(1, m.WEEKS + 1):
    c = col(7 + w)
    ws.cell(row=util_r, column=7 + w, value=f"={c}{r0}/{c}{cap_r}").number_format = "0%"
ws.cell(row=util_r + 1, column=2, value="Tải tối đa một vai trò (so với năng lực 1 người)").font = BOLD
for w in range(1, m.WEEKS + 1):
    c = col(7 + w)
    ws.cell(row=util_r + 1, column=7 + w, value=f"=MAX({c}{r0+1}:{c}{r0+4})/{G['cap']}").number_format = "0%"
ws.conditional_formatting.add(f"H{util_r+1}:{col(7+m.WEEKS)}{util_r+1}", CellIsRule(operator="greaterThan", formula=["1.15"], fill=PatternFill("solid", bgColor="EA9999")))
ws.conditional_formatting.add(f"H{util_r+1}:{col(7+m.WEEKS)}{util_r+1}", CellIsRule(operator="between", formula=["1.0001", "1.15"], fill=PatternFill("solid", bgColor="FFE599")))
pv_r = util_r + 3
ws.cell(row=pv_r, column=2, value="PV tuần (đ)").font = BOLD
ws.cell(row=pv_r + 1, column=2, value="PV tích luỹ (đ)").font = BOLD
for w in range(1, m.WEEKS + 1):
    c = col(7 + w)
    ws.cell(row=pv_r, column=7 + w, value=f"={c}{r0}*{G['rate']}").number_format = VND
    prev = f"+{col(6 + w)}{pv_r + 1}" if w > 1 else ""
    ws.cell(row=pv_r + 1, column=7 + w, value=f"={c}{pv_r}{prev}").number_format = VND
widths(ws, [9, 52, 7, 9, 8, 5, 5])
ws.freeze_panes = "H5"
PV_ROW = pv_r  # PV tuần ở hàng này, tích luỹ ở hàng sau
PV_TOTAL_ROW = r0

# ---------------------------------------------------------------- RuiRo
ws = wb.create_sheet("RuiRo")
title(ws, "Sổ rủi ro", "Điểm = Xác suất x Tác động (thang 1-5). Xác suất rời rạc: 1=10%, 2=30%, 3=50%, 4=70%, 5=90%. EMV = Tổn thất x Xác suất.")
header(ws, 4, ["Mã", "Rủi ro", "Nguyên nhân", "Hậu quả", "Nhóm", "P (1-5)", "I (1-5)", "Điểm", "Mức", "% xác suất", "Tổn thất nếu xảy ra (đ)",
               "EMV (đ)", "Chiến lược", "Kế hoạch ứng phó", "Dấu hiệu kích hoạt", "Chủ sở hữu", "Trạng thái"])
for i, rk in enumerate(m.RISKS, start=5):
    vals = [rk.rid, rk.name, rk.cause, rk.effect, rk.cat, rk.p, rk.i]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=6).fill = INPUT
    ws.cell(row=i, column=7).fill = INPUT
    ws.cell(row=i, column=8, value=f"=F{i}*G{i}")
    ws.cell(row=i, column=9, value=f'=IF(H{i}>=15,"Cao",IF(H{i}>=8,"Trung bình","Thấp"))')
    ws.cell(row=i, column=10, value=f"=CHOOSE(F{i},0.1,0.3,0.5,0.7,0.9)").number_format = "0%"
    ws.cell(row=i, column=11, value=rk.cost_if).number_format = VND
    ws.cell(row=i, column=11).fill = INPUT
    ws.cell(row=i, column=12, value=f"=K{i}*J{i}").number_format = VND
    for c, v in zip(range(13, 18), [rk.strategy, rk.response, rk.trigger, rk.owner, rk.status]):
        ws.cell(row=i, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for c in (8, 9, 10, 12):
        ws.cell(row=i, column=c).fill = CALC
RK_LAST = 4 + len(m.RISKS)
ws.cell(row=RK_LAST + 1, column=2, value="Tổng EMV").font = BOLD
ws.cell(row=RK_LAST + 1, column=12, value=f"=SUM(L5:L{RK_LAST})").number_format = VND
ws.cell(row=RK_LAST + 2, column=2, value="Dự phòng quản lý (đ)")
ws.cell(row=RK_LAST + 2, column=12, value=f"={RESERVE_TOTAL_CELL}-{BAC_CELL}").number_format = VND
ws.cell(row=RK_LAST + 3, column=2, value="Dự phòng so với EMV")
ws.cell(row=RK_LAST + 3, column=12, value=f"=L{RK_LAST+2}/L{RK_LAST+1}").number_format = "0%"
ws.conditional_formatting.add(f"I5:I{RK_LAST}", FormulaRule(formula=['I5="Cao"'], fill=PatternFill("solid", bgColor="EA9999")))
ws.conditional_formatting.add(f"I5:I{RK_LAST}", FormulaRule(formula=['I5="Trung bình"'], fill=PatternFill("solid", bgColor="FFE599")))
ws.conditional_formatting.add(f"I5:I{RK_LAST}", FormulaRule(formula=['I5="Thấp"'], fill=PatternFill("solid", bgColor="B6D7A8")))
box(ws, f"A4:Q{RK_LAST}")
mr = RK_LAST + 6
ws.cell(row=mr - 1, column=2, value="Ma trận rủi ro (số rủi ro trong ô). Hàng: Xác suất 5 xuống 1; cột: Tác động 1 đến 5").font = BOLD
for j in range(1, 6):
    c = ws.cell(row=mr, column=2 + j, value=f"I={j}")
    c.fill = SUB
    c.font = BOLD
for k, p in enumerate(range(5, 0, -1), start=1):
    ws.cell(row=mr + k, column=2, value=f"P={p}").fill = SUB
    for j in range(1, 6):
        cell = ws.cell(row=mr + k, column=2 + j, value=f"=COUNTIFS($F$5:$F${RK_LAST},{p},$G$5:$G${RK_LAST},{j})")
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
        score = p * j
        cell.fill = BAD if score >= 15 else (WARN if score >= 8 else GOOD)
widths(ws, [5, 30, 30, 32, 12, 7, 7, 7, 11, 9, 16, 14, 14, 54, 26, 16, 20])
ws.freeze_panes = "C5"

# ---------------------------------------------------------------- RACI
ws = wb.create_sheet("RACI")
title(ws, "Ma trận trách nhiệm RACI", "R: thực hiện, A: chịu trách nhiệm cuối (đúng một người), C: được tham vấn, I: được thông báo. A/R: vừa chịu trách nhiệm vừa thực hiện.")
header(ws, 4, ["Mã", "Hạng mục"] + [m.RACI_COL_NAMES[c] for c in m.RACI_COLS] + ["Số A", "Kiểm tra"])
for i, (code, name, cells) in enumerate(m.RACI, start=5):
    ws.cell(row=i, column=1, value=code)
    ws.cell(row=i, column=2, value=name)
    for j, c in enumerate(m.RACI_COLS, start=3):
        cell = ws.cell(row=i, column=j, value=cells[c])
        cell.alignment = Alignment(horizontal="center")
        cell.fill = INPUT
    ws.cell(row=i, column=9, value=f'=COUNTIF(C{i}:H{i},"A*")')
    ws.cell(row=i, column=10, value=f'=IF(I{i}=1,"OK","LỖI")')
RC_LAST = 4 + len(m.RACI)
ws.conditional_formatting.add(f"C5:H{RC_LAST}", FormulaRule(formula=['LEFT(C5,1)="A"'], fill=PatternFill("solid", bgColor="EA9999")))
ws.conditional_formatting.add(f"C5:H{RC_LAST}", FormulaRule(formula=['C5="R"'], fill=PatternFill("solid", bgColor="9FC5E8")))
ws.conditional_formatting.add(f"J5:J{RC_LAST}", FormulaRule(formula=['J5="LỖI"'], fill=PatternFill("solid", bgColor="EA9999")))
box(ws, f"A4:J{RC_LAST}")
ws.cell(row=RC_LAST + 2, column=2, value="Số hạng mục có đúng một A").font = BOLD
ws.cell(row=RC_LAST + 2, column=9, value=f'=COUNTIF(J5:J{RC_LAST},"OK")')
widths(ws, [9, 36, 16, 16, 16, 16, 18, 18, 7, 9])

# ---------------------------------------------------------------- TaiChinh
ws = wb.create_sheet("TaiChinh")
title(ws, "Phân tích tài chính", "Chi phí phát triển lấy từ sheet WBS. Lợi ích tính từ giả định ở sheet GiaDinh.")
rows = [
    ("Lượt mượn trả mỗi năm", f"={G['readers']}*{G['lpr']}", VND),
    ("Lợi ích 1: giờ thủ thư tiết kiệm khi mượn trả (đ/năm)", f"=B5*{G['min_loan']}/60*{G['lib_rate']}", VND),
    ("Lợi ích 2: giờ tra cứu hộ tiết kiệm (đ/năm)", f"={G['lookups']}*{G['min_look']}/60*{G['lib_rate']}", VND),
    ("Lợi ích 3: giờ làm báo cáo tiết kiệm (đ/năm)", f"={G['rep_h']}*{G['lib_rate']}", VND),
    ("Lợi ích 4: giảm thất thoát sách (đ/năm)", f"={G['lost_v']}*{G['lost_r']}", VND),
    ("Tổng lợi ích đầy đủ (đ/năm)", "=SUM(B6:B9)", VND),
    ("Vốn đầu tư ban đầu = BAC + dự phòng (đ)", f"={RESERVE_TOTAL_CELL}", VND),
    ("Chi phí vận hành mỗi năm (đ)", f"={G['host']}+{G['mh']}*{G['rate']}+{G['train']}", VND),
]
header(ws, 4, ["Khoản mục", "Giá trị"])
for i, (lab, f, fmt) in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=lab)
    c = ws.cell(row=i, column=2, value=f)
    c.number_format = fmt
    c.fill = CALC
ws["A10"].font = BOLD
ws["A14"] = "Dòng tiền"
ws["A14"].font = BOLD
header(ws, 15, ["Năm", "Lợi ích (đ)", "Chi phí (đ)", "Dòng tiền ròng (đ)", "Hệ số chiết khấu", "Giá trị hiện tại (đ)", "Luỹ kế không chiết khấu (đ)"])
ws.cell(row=16, column=1, value=0)
ws.cell(row=16, column=2, value=0)
ws.cell(row=16, column=3, value="=B11")
ws.cell(row=16, column=4, value="=B16-C16")
ws.cell(row=16, column=5, value=f"=1/(1+{G['disc']})^A16")
ws.cell(row=16, column=6, value="=D16*E16")
ws.cell(row=16, column=7, value="=D16")
for y, adkey in zip((1, 2, 3), ("ad1", "ad2", "ad3")):
    r = 16 + y
    ws.cell(row=r, column=1, value=y)
    ws.cell(row=r, column=2, value=f"=$B$10*{G[adkey]}")
    ws.cell(row=r, column=3, value="=$B$12")
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
    ws.cell(row=r, column=5, value=f"=1/(1+{G['disc']})^A{r}")
    ws.cell(row=r, column=6, value=f"=D{r}*E{r}")
    ws.cell(row=r, column=7, value=f"=G{r-1}+D{r}")
for r in range(16, 20):
    for c in (2, 3, 4, 6, 7):
        ws.cell(row=r, column=c).number_format = VND
        ws.cell(row=r, column=c).fill = CALC
    ws.cell(row=r, column=5).number_format = "0.0000"
    ws.cell(row=r, column=5).fill = CALC
box(ws, "A15:G19")
ws["A21"] = "NPV (3 năm)"
ws["B21"] = "=SUM(F16:F19)"
ws["A22"] = "ROI (3 năm) = (tổng dòng tiền ròng các năm 1-3 - vốn) / vốn"
ws["B22"] = "=(SUM(D17:D19)-B11)/B11"
ws["A23"] = "Thời gian hoàn vốn (năm, không chiết khấu)"
ws["B23"] = "=IF(D17>=B11,B11/D17,IF(D17+D18>=B11,1+(B11-D17)/D18,2+(B11-D17-D18)/D19))"
ws["A24"] = "Lợi ích / chi phí (BCR) = tổng PV lợi ích / (vốn + PV vận hành)"
ws["B24"] = f"=(B17/(1+{G['disc']})+B18/(1+{G['disc']})^2+B19/(1+{G['disc']})^3)/(B11+C17/(1+{G['disc']})+C18/(1+{G['disc']})^2+C19/(1+{G['disc']})^3)"
for r, fmt in ((21, VND), (22, "0.0%"), (23, "0.00"), (24, "0.00")):
    ws.cell(row=r, column=2).number_format = fmt
    ws.cell(row=r, column=2).fill = CALC
    ws.cell(row=r, column=2).font = BOLD
ws["A27"] = "Phân tích độ nhạy: NPV (đ) theo mức lợi ích và tỉ lệ chiết khấu"
ws["A27"].font = BOLD
header(ws, 28, ["Lợi ích so với kế hoạch \\ Chiết khấu", 0.08, 0.10, 0.12, 0.15])
for k, bf in enumerate((1.0, 0.7, 0.5, 0.3), start=29):
    ws.cell(row=k, column=1, value=bf).number_format = "0%"
    for j in range(2, 6):
        dc = f"{col(j)}$28"
        ws.cell(row=k, column=j, value=(
            f"=-$B$11+($B$10*{G['ad1']}*$A{k}-$B$12)/(1+{dc})+($B$10*{G['ad2']}*$A{k}-$B$12)/(1+{dc})^2"
            f"+($B$10*{G['ad3']}*$A{k}-$B$12)/(1+{dc})^3")).number_format = VND
        ws.cell(row=k, column=j).fill = CALC
box(ws, "A28:E32")
ws["A34"] = "Điểm hoà vốn: mức lợi ích tối thiểu để NPV = 0 (so với kế hoạch)"
ws["B34"] = (f"=($B$11+$B$12*(1/(1+{G['disc']})+1/(1+{G['disc']})^2+1/(1+{G['disc']})^3))/"
             f"($B$10*({G['ad1']}/(1+{G['disc']})+{G['ad2']}/(1+{G['disc']})^2+{G['ad3']}/(1+{G['disc']})^3))")
ws["B34"].number_format = "0.0%"
ws["B34"].fill = CALC
widths(ws, [70, 22, 22, 22, 18, 22, 28])

# ---------------------------------------------------------------- EVA
ws = wb.create_sheet("EVA")
title(ws, "Phân tích giá trị thu được (EVA)",
      "PV từ lịch kế hoạch. EV = EV trước + giờ kế hoạch tuần x hiệu suất tiến độ (chặn ở BAC). AC = AC trước + (EV tuần) x hệ số chi phí. Số mô phỏng có kiểm soát.")
header(ws, 4, ["Tuần", "Từ ngày", "PV tuần (đ)", "Hiệu suất tiến độ", "Hệ số chi phí", "EV tuần (đ)", "AC tuần (đ)", "PV (đ)", "EV (đ)", "AC (đ)",
               "SV (đ)", "CV (đ)", "SPI", "CPI", "EAC = BAC/CPI (đ)", "EAC theo SPI x CPI (đ)", "ETC (đ)", "VAC (đ)", "TCPI", "Mốc báo cáo"])
EV_FIRST = 5
phase_of = lambda w: 1 if w <= 4 else (2 if w <= 8 else (3 if w <= 12 else 4))
for w in range(1, m.WEEKS + 1):
    r = EV_FIRST + w - 1
    s, _ = m.week_dates(w)
    ws.cell(row=r, column=1, value=w)
    ws.cell(row=r, column=2, value=s).number_format = "dd/mm/yyyy"
    pvcol = col(7 + w)
    ws.cell(row=r, column=3, value=f"=NguonLuc!{pvcol}{PV_ROW}").number_format = VND
    ws.cell(row=r, column=4, value=f"=GiaDinh!$B${EVA_PHASE_ROW[phase_of(w)]}").number_format = "0.00"
    ws.cell(row=r, column=5, value=f"=GiaDinh!$C${EVA_PHASE_ROW[phase_of(w)]}").number_format = "0.00"
    prev_ev = f"I{r-1}" if w > 1 else "0"
    prev_ac = f"J{r-1}" if w > 1 else "0"
    if w == m.WEEKS:
        ws.cell(row=r, column=6, value=f"={BAC_CELL}-{prev_ev}")  # tuần cuối bù để hoàn thành 100%
    else:
        ws.cell(row=r, column=6, value=f"=MIN(C{r}*D{r},{BAC_CELL}-{prev_ev})")
    ws.cell(row=r, column=7, value=f"=F{r}*E{r}")
    ws.cell(row=r, column=8, value=f"=NguonLuc!{pvcol}{PV_ROW + 1}")
    ws.cell(row=r, column=9, value=f"={prev_ev}+F{r}")
    ws.cell(row=r, column=10, value=f"={prev_ac}+G{r}")
    ws.cell(row=r, column=11, value=f"=I{r}-H{r}")
    ws.cell(row=r, column=12, value=f"=I{r}-J{r}")
    ws.cell(row=r, column=13, value=f"=I{r}/H{r}")
    ws.cell(row=r, column=14, value=f"=I{r}/J{r}")
    ws.cell(row=r, column=15, value=f"={BAC_CELL}/N{r}")
    ws.cell(row=r, column=16, value=f"=J{r}+({BAC_CELL}-I{r})/(N{r}*M{r})")
    ws.cell(row=r, column=17, value=f"=O{r}-J{r}")
    ws.cell(row=r, column=18, value=f"={BAC_CELL}-O{r}")
    ws.cell(row=r, column=19, value=f'=IF({BAC_CELL}-J{r}=0,"",({BAC_CELL}-I{r})/({BAC_CELL}-J{r}))')
    ws.cell(row=r, column=20, value="Báo cáo" if w in m.EVA_REPORT_WEEKS else "")
    for c in (3, 6, 7, 8, 9, 10, 11, 12, 15, 16, 17, 18):
        ws.cell(row=r, column=c).number_format = VND
    for c in (13, 14, 19):
        ws.cell(row=r, column=c).number_format = "0.000"
    for c in range(3, 21):
        ws.cell(row=r, column=c).fill = CALC
    ws.cell(row=r, column=4).fill = CALC
    ws.cell(row=r, column=5).fill = CALC
EV_LAST = EV_FIRST + m.WEEKS - 1
ws.conditional_formatting.add(f"A{EV_FIRST}:T{EV_LAST}", FormulaRule(formula=['$T5="Báo cáo"'], font=Font(bold=True), fill=PatternFill("solid", bgColor="DDEBF7")))
ws.conditional_formatting.add(f"M{EV_FIRST}:N{EV_LAST}", CellIsRule(operator="lessThan", formula=["0.85"], fill=PatternFill("solid", bgColor="EA9999")))
ws.conditional_formatting.add(f"M{EV_FIRST}:N{EV_LAST}", CellIsRule(operator="between", formula=["0.85", "0.9499"], fill=PatternFill("solid", bgColor="FFE599")))
ws.conditional_formatting.add(f"M{EV_FIRST}:N{EV_LAST}", CellIsRule(operator="greaterThanOrEqual", formula=["0.95"], fill=PatternFill("solid", bgColor="B6D7A8")))
box(ws, f"A4:T{EV_LAST}")
ws.cell(row=EV_LAST + 2, column=1, value="BAC (đ)").font = BOLD
ws.cell(row=EV_LAST + 2, column=3, value=f"={BAC_CELL}").number_format = VND
ws.cell(row=EV_LAST + 3, column=1, value="Ngưỡng: SPI, CPI >= 0,95 xanh; 0,85-0,95 vàng; < 0,85 đỏ").font = Font(italic=True)
chart = LineChart()
chart.title = "PV, EV, AC tích luỹ (đ)"
chart.height = 9
chart.width = 22
chart.y_axis.title = "đồng"
chart.x_axis.title = "tuần"
data = Reference(ws, min_col=8, max_col=10, min_row=4, max_row=EV_LAST)
chart.add_data(data, titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=EV_FIRST, max_row=EV_LAST))
ws.add_chart(chart, f"A{EV_LAST + 6}")
widths(ws, [6, 11, 14, 10, 10, 14, 14, 14, 14, 14, 13, 13, 8, 8, 16, 17, 14, 14, 8, 10])
ws.freeze_panes = "C5"

# ---------------------------------------------------------------- Timelog
ws = wb.create_sheet("Timelog")
title(ws, "Giờ công thật ghi nhận của nhóm", "Nguồn: docs/pm/timelog.csv. Đây là số đo thật, không dùng trực tiếp cho EVA (xem báo cáo, mục EVA).")
header(ws, 4, ["Ngày", "Người", "Mã WBS", "Giờ", "Ghi chú"])
log = Path(__file__).resolve().parents[1] / "pm" / "timelog.csv"
TL_FIRST = 5
n = 0
with open(log, encoding="utf-8", newline="") as fh:
    for i, row in enumerate(csv.DictReader(fh), start=TL_FIRST):
        ws.cell(row=i, column=1, value=row["date"])
        ws.cell(row=i, column=2, value=row["person"])
        ws.cell(row=i, column=3, value=row["wbs_id"])
        ws.cell(row=i, column=4, value=float(row["hours"]))
        ws.cell(row=i, column=5, value=row["note"])
        n += 1
TL_LAST = TL_FIRST + n - 1
ws.cell(row=TL_LAST + 1, column=3, value="Tổng").font = BOLD
ws.cell(row=TL_LAST + 1, column=4, value=f"=SUM(D{TL_FIRST}:D{TL_LAST})").font = BOLD
ws.cell(row=TL_LAST + 3, column=1, value="Tổng giờ thật theo pha (mã WBS bắt đầu bằng chữ số pha)").font = BOLD
for k, (code, name) in enumerate(m.WBS_L1.items()):
    r = TL_LAST + 4 + k
    ws.cell(row=r, column=1, value=f"{code}. {name}")
    ws.cell(row=r, column=4, value=f'=SUMPRODUCT((LEFT($C${TL_FIRST}:$C${TL_LAST},1)="{code}")*$D${TL_FIRST}:$D${TL_LAST})')
widths(ws, [14, 8, 9, 7, 110])

# ---------------------------------------------------------------- ChatLuong
ws = wb.create_sheet("ChatLuong")
title(ws, "Mốc chất lượng và tiêu chí chấp nhận", "Kết quả thật lấy từ pipeline CI của dự án (run #25 trên main).")
header(ws, 4, ["Loại kiểm thử", "Phạm vi", "Công cụ", "Tiêu chí chấp nhận", "Kết quả thật", "Đạt?"])
qa_rows = [
    ("Kiểm thử đơn vị", "Service, guard, client của Backend", "Jest", "Coverage dòng >= 80% trên lớp nghiệp vụ", "119 test; dòng 98,75%, câu lệnh 97,51%, hàm 98,49%, nhánh 81,9%", "Đạt"),
    ("Kiểm thử đơn vị AI", "Mô hình gợi ý, dự báo, API", "pytest", "Coverage >= 80%", "26 test; coverage 93,6%", "Đạt"),
    ("Kiểm thử tích hợp (e2e)", "API trên PostgreSQL thật", "Jest + Supertest", "Luồng mượn, trả, phạt đúng; mã lỗi đúng", "15 test, 3 suite", "Đạt"),
    ("Kiểm thử hệ thống", "Gọi API thật qua gateway", "Script smoke (Python)", "Mọi luồng chính trả mã đúng", "100 lượt gọi xanh (3 sprint)", "Đạt"),
    ("Kiểm thử cài đặt", "docker compose up trên máy sạch", "GitHub Actions (compose-smoke)", "3 địa chỉ /health trả 200 qua gateway", "Xanh trên runner sạch", "Đạt"),
    ("Kiểm thử chấp nhận (UAT)", "Kịch bản thủ thư và bạn đọc trên bản chạy thật", "Thủ công theo kịch bản", "Hoàn tất 100% kịch bản, 0 lỗi nghiêm trọng", "Đã chạy các luồng chính trên Vercel, Render, Neon", "Đạt"),
    ("Chỉ số mô hình AI", "Gợi ý sách và dự báo trả trễ", "scikit-learn", "Precision@5 >= 0,15; ROC-AUC >= 0,70", "0,309 và 0,76", "Đạt"),
]
for i, row in enumerate(qa_rows, start=5):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
box(ws, f"A4:F{4+len(qa_rows)}")
widths(ws, [26, 40, 28, 44, 52, 8])

out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("ULS_PM_Plan.xlsx")
for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 90
wb.save(out)
print("wrote", out)
