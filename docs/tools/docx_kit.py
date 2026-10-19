"""Bộ dựng nội dung cho hai mẫu Word của giảng viên (giữ nguyên bìa, đầu/chân trang, kiểu chữ của mẫu).

Cách làm: chụp các đoạn mẫu (tiêu đề, thân bài, gạch đầu dòng) trước khi xoá phần thân mẫu, rồi nhân bản chúng
cho nội dung thật. Bảng, hình, mục lục và phần xoay ngang được dựng bằng python-docx.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

FONT = "Times New Roman"
CODE_FONT = "Consolas"


# --------------------------------------------------------------------------- đọc bảng từ báo cáo Markdown
def clean_md(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).strip()


def split_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [clean_md(c) for c in cells]


class MdSource:
    """Lấy lại các bảng đã có trong báo cáo QLDAPM để con số khớp tuyệt đối với báo cáo gốc."""

    def __init__(self, path: Path):
        self.lines = path.read_text(encoding="utf-8").splitlines()

    def table(self, caption_prefix: str) -> tuple[list[str], list[list[str]]]:
        for i, line in enumerate(self.lines):
            if line.startswith(caption_prefix):
                break
        else:
            raise LookupError(caption_prefix)
        j = i - 1
        while j >= 0 and not self.lines[j].startswith("|"):
            j -= 1
        end = j
        while j >= 0 and self.lines[j].startswith("|"):
            j -= 1
        block = self.lines[j + 1 : end + 1]
        header = split_row(block[0])
        rows = [split_row(r) for r in block[2:]]
        return header, rows


# --------------------------------------------------------------------------- đoạn văn
_CODE_AUTO = re.compile(
    r"""(?:(?:GET|POST|PUT|PATCH|DELETE)\s+/[\w/{}.\-]*)
    |(?:/api/[\w/{}.\-]+)
    |(?:\b[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+\b)
    |(?:\b(?:ACTIVE|BORROWED|AVAILABLE|RETURNED|LOST|OVERDUE|HELD|FULFILLED|null)\b)
    |(?:\b(?:docs|backend|frontend)/[\w./\-]+)
    |(?:\b[\w\-]+\.(?:xlsx|docx|md|py|csv|sql|yml)\b)
    |(?:/health\b)
    |(?:docker compose(?: up)?)""",
    re.X,
)


def _split_code(chunk: str) -> list[tuple[str, bool]]:
    """Tách một đoạn chữ thành (chữ, là_mã): `...` là mã rõ ràng, phần còn lại tự nhận ra lệnh, mã lỗi, đường dẫn."""
    out: list[tuple[str, bool]] = []
    for piece in re.split(r"(`[^`]+`)", chunk):
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            out.append((piece[1:-1], True))
            continue
        pos = 0
        for m in _CODE_AUTO.finditer(piece):
            if m.start() > pos:
                out.append((piece[pos : m.start()], False))
            out.append((m.group(0), True))
            pos = m.end()
        if pos < len(piece):
            out.append((piece[pos:], False))
    return out


def _runs_with_bold(text: str) -> list[tuple[str, bool, bool]]:
    """(chữ, đậm, mã) theo thứ tự xuất hiện; `**đậm**` và mã được tách riêng."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        for chunk, is_code in _split_code(part[2:-2] if is_bold else part):
            out.append((chunk, is_bold, is_code))
    return out


class Template:
    """Một mẫu Word đang được điền."""

    def __init__(self, path: Path):
        self.doc = Document(path)
        self.body = self.doc.element.body
        self.protos: dict[str, "copy.Any"] = {}
        self.default_half_pt: int | None = None  # cỡ chữ mặc định (nửa điểm) cho đoạn tạo mới
        self.force_font_pt: float | None = None  # ép cỡ chữ cho mọi ô bảng

    # ---- đoạn mẫu
    def capture(self, name: str, paragraph: Paragraph, strip_bold: bool = False) -> None:
        proto = copy.deepcopy(paragraph._p)
        if strip_bold:  # đoạn mẫu có cụm đầu in đậm: bỏ đậm để chỉ đậm chỗ đánh dấu **
            for tag in ("w:b", "w:bCs"):
                for node in proto.iter(qn(tag)):
                    node.getparent().remove(node)
        self.protos[name] = proto

    def _base_rpr(self, proto) -> "OxmlElement | None":
        for r in proto.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                return rpr
        ppr = proto.find(qn("w:pPr"))
        if ppr is not None:
            return ppr.find(qn("w:rPr"))
        return None

    def make_para(self, kind: str, text: str, *, bold: bool | None = None, italic: bool | None = None,
                  outline: int | None = None, size_half_pt: int | None = None, align: str | None = None,
                  page_break_before: bool = False, left_twips: int | None = None, keep_next: bool = False,
                  num_id: int | None = None, list_num: int | None = None, list_level: int = 0):
        """Nhân bản đoạn mẫu `kind` rồi đặt lại phần chữ. `**đậm**` được hiểu là chữ đậm."""
        proto = self.protos[kind]
        el = copy.deepcopy(proto)
        base = self._base_rpr(proto)
        for child in list(el):
            if child.tag != qn("w:pPr"):
                el.remove(child)
        ppr = el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            el.insert(0, ppr)
        if outline is not None:
            for old in ppr.findall(qn("w:outlineLvl")):
                ppr.remove(old)
            node = OxmlElement("w:outlineLvl")
            node.set(qn("w:val"), str(outline))
            _insert_ppr_child(ppr, node)
        if page_break_before:
            node = OxmlElement("w:pageBreakBefore")
            _insert_ppr_child(ppr, node)
        if list_num is not None:
            for old in ppr.findall(qn("w:numPr")) + ppr.findall(qn("w:ind")):
                ppr.remove(old)
            num_pr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), str(list_level))
            num_id_el = OxmlElement("w:numId")
            num_id_el.set(qn("w:val"), str(list_num))
            num_pr.append(ilvl)
            num_pr.append(num_id_el)
            _insert_ppr_child(ppr, num_pr)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(360 * (list_level + 1)))
            ind.set(qn("w:hanging"), "360")
            _insert_ppr_child(ppr, ind)
        if num_id is not None:
            num = ppr.find(qn("w:numPr"))
            if num is not None and num.find(qn("w:numId")) is not None:
                num.find(qn("w:numId")).set(qn("w:val"), str(num_id))
        if keep_next and ppr.find(qn("w:keepNext")) is None:
            _insert_ppr_child(ppr, OxmlElement("w:keepNext"))
        if align:
            for old in ppr.findall(qn("w:jc")):
                ppr.remove(old)
            node = OxmlElement("w:jc")
            node.set(qn("w:val"), align)
            _insert_ppr_child(ppr, node)
        if left_twips is not None:
            for old in ppr.findall(qn("w:ind")):
                ppr.remove(old)
            node = OxmlElement("w:ind")
            node.set(qn("w:left"), str(left_twips))
            node.set(qn("w:hanging"), "360")
            _insert_ppr_child(ppr, node)
        for chunk, is_bold, is_code in _runs_with_bold(text):
            r = OxmlElement("w:r")
            rpr = copy.deepcopy(base) if base is not None else OxmlElement("w:rPr")
            if is_code:
                _set_fonts(rpr, CODE_FONT)
            for tag in ("w:lastRenderedPageBreak",):
                for x in rpr.findall(qn(tag)):
                    rpr.remove(x)
            if bold is not None or is_bold:  # không có chỉ định thì giữ nguyên độ đậm của đoạn mẫu
                want_bold = bool(bold) or is_bold
                _set_flag(rpr, "w:b", want_bold)
                _set_flag(rpr, "w:bCs", want_bold)
            if italic is not None:
                _set_flag(rpr, "w:i", italic)
                _set_flag(rpr, "w:iCs", italic)
            size = size_half_pt or self.default_half_pt
            if size:
                _set_val(rpr, "w:sz", size)
                _set_val(rpr, "w:szCs", size)
            r.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = chunk
            r.append(t)
            el.append(r)
        return el

    def add(self, kind: str, text: str, **kw):
        el = self.make_para(kind, text, **kw)
        self._append(el)
        return el

    def make_dash_list(self) -> int:
        """Danh sách gạch đầu dòng nhiều cấp: cấp 1 dấu -, cấp 2 dấu +, cấp 3 dấu chấm (•)."""
        numbering = self.doc.part.numbering_part.element
        abstracts = numbering.findall(qn("w:abstractNum"))
        new_abs = max(int(a.get(qn("w:abstractNumId"))) for a in abstracts) + 1
        absn = OxmlElement("w:abstractNum")
        absn.set(qn("w:abstractNumId"), str(new_abs))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "hybridMultilevel")
        absn.append(multi)
        for lvl, glyph in enumerate(["-", "+", "\u2022"]):
            node = OxmlElement("w:lvl")
            node.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", glyph), ("w:lvlJc", "left")):
                child = OxmlElement(tag)
                child.set(qn("w:val"), val)
                node.append(child)
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(360 * (lvl + 1)))
            ind.set(qn("w:hanging"), "360")
            ppr.append(ind)
            node.append(ppr)
            rpr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            for att in ("w:ascii", "w:hAnsi", "w:cs"):
                fonts.set(qn(att), FONT)
            rpr.append(fonts)
            node.append(rpr)
            absn.append(node)
        abstracts[-1].addnext(absn)
        nums = numbering.findall(qn("w:num"))
        new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        ab = OxmlElement("w:abstractNumId")
        ab.set(qn("w:val"), str(new_abs))
        num.append(ab)
        nums[-1].addnext(num)
        return new_id

    def new_list_num(self, base_num_id: int) -> int:
        """Tạo bản sao đánh số của danh sách `base_num_id` bắt đầu lại từ 1 (mỗi nhóm gạch đầu dòng a), b), c)... riêng)."""
        numbering = self.doc.part.numbering_part.element
        nums = numbering.findall(qn("w:num"))
        base = next(n for n in nums if n.get(qn("w:numId")) == str(base_num_id))
        abstract = base.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        ab = OxmlElement("w:abstractNumId")
        ab.set(qn("w:val"), abstract)
        num.append(ab)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        st = OxmlElement("w:startOverride")
        st.set(qn("w:val"), "1")
        ov.append(st)
        num.append(ov)
        nums[-1].addnext(num)
        return new_id

    def _append(self, el) -> None:
        sect = self.body.find(qn("w:sectPr"))
        sect.addprevious(el)

    def page_break(self) -> None:
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        r.append(br)
        p.append(r)
        self._append(p)

    def clear_after(self, paragraph: Paragraph) -> None:
        """Xoá mọi phần tử thân bài sau `paragraph`, giữ lại sectPr cuối."""
        node = paragraph._p.getnext()
        while node is not None and node.tag != qn("w:sectPr"):
            nxt = node.getnext()
            self.body.remove(node)
            node = nxt

    # ---- hình
    def picture(self, path: Path, max_w_cm: float, caption: str | None = None, caption_kind: str = "caption",
                max_h_cm: float = 20.0, caption_kw: dict | None = None) -> None:
        from PIL import Image

        with Image.open(path) as im:
            w_px, h_px = im.size
        width = min(max_w_cm, max_h_cm * w_px / h_px)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = bool(caption)
        p.add_run().add_picture(str(path), width=Cm(width))
        if caption:
            self.add(caption_kind, caption, **(caption_kw or {}))

    # ---- mục lục
    def toc_paragraph(self, levels: str = "1-2"):
        """Trường mục lục; Word cập nhật số trang khi mở hoặc khi chạy update_fields."""
        p = OxmlElement("w:p")

        def run_with(child):
            r = OxmlElement("w:r")
            r.append(child)
            p.append(r)

        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        run_with(begin)
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f' TOC \\o "{levels}" \\h \\z \\u '
        run_with(instr)
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        run_with(sep)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "Nhấn chuột phải rồi chọn Update Field để cập nhật mục lục."
        r.append(t)
        p.append(r)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run_with(end)
        return p

    # ---- phần xoay ngang
    def begin_landscape(self) -> None:
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        _drop_page_number_restart(sec)
        sec.orientation = WD_ORIENT.LANDSCAPE
        w, h = sec.page_width, sec.page_height
        if w < h:
            sec.page_width, sec.page_height = h, w

    def end_landscape(self) -> None:
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        _drop_page_number_restart(sec)
        sec.orientation = WD_ORIENT.PORTRAIT
        w, h = sec.page_width, sec.page_height
        if w > h:
            sec.page_width, sec.page_height = h, w

    # ---- bảng
    def table(self, header: list[str], rows: list[list[str]], widths_cm: list[float], *, style: str,
              font_pt: float = 11, header_fill: str | None = "E7E9F0", bold_first_col: bool = False,
              center_cols: set[int] | None = None, total_row: bool = False, keep_header: bool = True,
              gap: bool = True, keep_last: bool = False):
        center_cols = center_cols or set()
        if self.force_font_pt:
            font_pt = self.force_font_pt
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.style = self.doc.styles[style]
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        for old in tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        total_tw = int(sum(widths_cm) / 2.54 * 1440)
        tblw = OxmlElement("w:tblW")
        tblw.set(qn("w:w"), str(total_tw))
        tblw.set(qn("w:type"), "dxa")
        tblPr.append(tblw)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tblPr.append(jc)
        order = {qn(t): i for i, t in enumerate(_TBLPR_ORDER)}
        for child in sorted(list(tblPr), key=lambda c: order.get(c.tag, 99)):
            tblPr.remove(child)
            tblPr.append(child)
        grid = tbl._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))
        all_rows = [header] + rows
        for ri, (tr, values) in enumerate(zip(tbl.rows, all_rows)):
            trPr = tr._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            trPr.append(cant)
            if ri == 0 and keep_header:
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            last = total_row and ri == len(all_rows) - 1
            for ci, (cell, value) in enumerate(zip(tr.cells, values)):
                cell.width = Cm(widths_cm[ci])
                if ri == 0 and header_fill:
                    _shade(cell, header_fill)
                lines = str(value).split("\n")
                for li, line in enumerate(lines):
                    para = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    pf = para.paragraph_format
                    pf.keep_with_next = (ri == 0) or (len(all_rows) <= 5 and ri < len(all_rows) - 1) or (keep_last and ri == len(all_rows) - 1)
                    pf.space_before = Pt(2)
                    pf.space_after = Pt(2)
                    pf.line_spacing = 1.0
                    if ri == 0 or ci in center_cols:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for chunk, is_bold, is_code in _runs_with_bold(line):
                        run = para.add_run(chunk)
                        run.font.name = CODE_FONT if is_code else FONT
                        run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CODE_FONT if is_code else FONT)
                        run._r.get_or_add_rPr().rFonts.set(qn("w:cs"), CODE_FONT if is_code else FONT)
                        run.font.size = Pt(font_pt)
                        run.font.bold = True if (ri == 0 or is_bold or last or (bold_first_col and ci == 0)) else False
        if gap:  # đoạn trống nhỏ sau bảng để các khối không dính nhau
            spacer = self.doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(6)
        return tbl


# --------------------------------------------------------------------------- tiện ích XML
_PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr",
    "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
    "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
]


def _insert_ppr_child(ppr, node) -> None:
    """Chèn đúng thứ tự phần tử của w:pPr để Word không báo tệp hỏng."""
    order = {qn(tag): i for i, tag in enumerate(_PPR_ORDER)}
    idx = order.get(node.tag, len(_PPR_ORDER))
    for i, child in enumerate(list(ppr)):
        if order.get(child.tag, len(_PPR_ORDER)) > idx:
            ppr.insert(i, node)
            return
    ppr.append(node)


_TBLPR_ORDER = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
                "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                "w:tblLayout", "w:tblCellMar", "w:tblLook"]

_RPR_ORDER = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
              "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
              "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz",
              "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
              "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath"]


def _insert_rpr_child(rpr, node) -> None:
    order = {qn(tag): i for i, tag in enumerate(_RPR_ORDER)}
    idx = order.get(node.tag, len(_RPR_ORDER))
    for i, child in enumerate(list(rpr)):
        if order.get(child.tag, len(_RPR_ORDER)) > idx:
            rpr.insert(i, node)
            return
    rpr.append(node)


def _set_fonts(rpr, name: str) -> None:
    for old in rpr.findall(qn("w:rFonts")):
        rpr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(att), name)
    _insert_rpr_child(rpr, fonts)


def _set_flag(rpr, tag: str, on: bool) -> None:
    for old in rpr.findall(qn(tag)):
        rpr.remove(old)
    if on:
        _insert_rpr_child(rpr, OxmlElement(tag))


def _set_val(rpr, tag: str, val: int) -> None:
    for old in rpr.findall(qn(tag)):
        rpr.remove(old)
    node = OxmlElement(tag)
    node.set(qn("w:val"), str(val))
    _insert_rpr_child(rpr, node)


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _drop_page_number_restart(section) -> None:
    sectPr = section._sectPr
    for node in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(node)


# --------------------------------------------------------------------------- bìa
def rebuild_label_line(paragraph: Paragraph, label: str, value: str, *, tab_twips: int = 3700,
                       base_indent: int = 0, size_half_pt: int = 28, font: str = "Arial") -> None:
    """Dòng bìa dạng "Nhãn:  giá trị", giá trị thẳng hàng nhờ một tab và thụt treo."""
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    ppr = p.find(qn("w:pPr"))
    for tag in ("w:ind", "w:tabs"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:pos"), str(base_indent + tab_twips))
    tabs.append(tab)
    _insert_ppr_child(ppr, tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(base_indent + tab_twips))
    ind.set(qn("w:hanging"), str(tab_twips))
    _insert_ppr_child(ppr, ind)

    def run(text: str, bold: bool, with_tab: bool = False):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for att in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(att), font)
        rpr.append(fonts)
        if bold:
            rpr.append(OxmlElement("w:b"))
            rpr.append(OxmlElement("w:bCs"))
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0D0D0D")
        rpr.append(color)
        for tag in ("w:sz", "w:szCs"):
            node = OxmlElement(tag)
            node.set(qn("w:val"), str(size_half_pt))
            rpr.append(node)
        r.append(rpr)
        if with_tab:
            r.append(OxmlElement("w:tab"))
        if text:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r.append(t)
        p.append(r)

    run(label, True)
    run(value, False, with_tab=True)


# --------------------------------------------------------------------------- một phông chữ cho cả tài liệu
def force_font(doc, name: str = FONT) -> None:
    """Đổi mọi phông trong nội dung, kiểu, đầu trang, chân trang (kể cả hộp văn bản ở bìa) sang `name`."""
    roots = [doc.element, doc.styles.element]
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
            roots.append(rel.target_part.element)
    for root in roots:
        for fonts in root.iter(qn("w:rFonts")):
            if fonts.get(qn("w:ascii")) == CODE_FONT:
                continue  # chữ thể hiện mã giữ phông Consolas
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme", "w:hint"):
                if fonts.get(qn(att)) is not None:
                    del fonts.attrib[qn(att)]
            for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(att), name)
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    if rpr.find(qn("w:rFonts")) is None:
        fonts = OxmlElement("w:rFonts")
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(att), name)
        rpr.insert(0, fonts)
